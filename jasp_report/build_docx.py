"""Builds GoodBelly_JASP_Final_Report.docx from results_jasp.json and fig/*.png (run analysis_jasp.py first)."""
import os, json
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__)); FIG = os.path.join(HERE, 'fig')
R = json.load(open(os.path.join(HERE, 'results_jasp.json')))
m1, ml, m2 = R['m1'], R['ml'], R['m2']
NAVY = RGBColor(0x1F, 0x4E, 0x79)

# ---------- number formats
def p_(p): return '<0.001' if p < 0.001 else f'{p:.3f}'
def n0(x): return f'{x:,.0f}'
def s0(x): return f'{x:+,.0f}'
def n1(x): return f'{x:,.1f}'
def n2(x): return f'{x:,.2f}'
def n3(x): return f'{x:.3f}'
def sg(x, k=2): return f'{x:+.{k}f}'

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
for a in ('left_margin', 'right_margin'): setattr(sec, a, Inches(0.9))
sec.top_margin, sec.bottom_margin = Inches(0.85), Inches(0.8)
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
st.paragraph_format.space_after = Pt(5); st.paragraph_format.line_spacing = 1.07
for name, size, before, after in [('Heading 1', 15, 14, 5), ('Heading 2', 12, 10, 3), ('Heading 3', 10.5, 7, 2)]:
    h = doc.styles[name]; h.font.name = 'Calibri'; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = NAVY
    h.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri'); h.element.rPr.rFonts.set(qn('w:ascii'), 'Calibri'); h.element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    h.paragraph_format.space_before = Pt(before); h.paragraph_format.space_after = Pt(after); h.paragraph_format.keep_with_next = True

def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hex_); tcPr.append(sh)

def borders(tbl, color='BFBFBF', sz='4'):
    tblPr = tbl._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), sz); x.set(qn('w:space'), '0'); x.set(qn('w:color'), color); b.append(x)
    anchor = next((x for x in (tblPr.find(qn('w:tblLayout')), tblPr.find(qn('w:tblCellMar')), tblPr.find(qn('w:tblLook'))) if x is not None), None)
    anchor.addprevious(b) if anchor is not None else tblPr.append(b)

def cell_margins(tbl, top=30, bottom=30, left=70, right=70):
    tblPr = tbl._tbl.tblPr; m = OxmlElement('w:tblCellMar')
    for k, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        x = OxmlElement(f'w:{k}'); x.set(qn('w:w'), str(v)); x.set(qn('w:type'), 'dxa'); m.append(x)
    look = tblPr.find(qn('w:tblLook')); look.addprevious(m) if look is not None else tblPr.append(m)

def runs(par, text, size=None, color=None, bold=None, italic=None):
    """text with **bold** segments"""
    for i, seg in enumerate(text.split('**')):
        if not seg: continue
        r = par.add_run(seg); r.bold = (i % 2 == 1) or bool(bold)
        if italic: r.italic = True
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
    return par

def P(text, size=None, after=None, align=None, italic=None, keep=False, color=None):
    par = doc.add_paragraph(); runs(par, text, size, color, italic=italic)
    if after is not None: par.paragraph_format.space_after = Pt(after)
    if align == 'c': par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if keep: par.paragraph_format.keep_with_next = True
    return par

def B(text, size=None, level=0):
    par = doc.add_paragraph(style='List Bullet'); runs(par, text, size)
    par.paragraph_format.space_after = Pt(2.5); par.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
    return par

def H(text, lvl=1): return doc.add_heading(text, lvl)

def table(header, rows, widths, align=None, size=8.5, bold_rows=(), shade_rows=None, header_fill='1F4E79'):
    t = doc.add_table(rows=1 + len(rows), cols=len(header)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    borders(t); cell_margins(t)
    align = align or ['l'] + ['r'] * (len(header) - 1)
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]; shade(c, header_fill); c.width = Inches(widths[j]); par = c.paragraphs[0]
        par.paragraph_format.space_after = Pt(0); r = par.add_run(h); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(255, 255, 255)
        par.alignment = {'l': WD_ALIGN_PARAGRAPH.LEFT, 'r': WD_ALIGN_PARAGRAPH.RIGHT, 'c': WD_ALIGN_PARAGRAPH.CENTER}[align[j]]
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.rows[i + 1].cells[j]; c.width = Inches(widths[j]); par = c.paragraphs[0]; par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            runs(par, str(v), size, bold=(i in bold_rows))
            par.alignment = {'l': WD_ALIGN_PARAGRAPH.LEFT, 'r': WD_ALIGN_PARAGRAPH.RIGHT, 'c': WD_ALIGN_PARAGRAPH.CENTER}[align[j]]
            if shade_rows and i in shade_rows: shade(c, shade_rows[i])
    for j, w_ in enumerate(widths): t.columns[j].width = Inches(w_)
    for r_ in t.rows[:-1]:
        for c_ in r_.cells:
            for q_ in c_.paragraphs: q_.paragraph_format.keep_with_next = True
    # repeat header, keep rows together
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); trPr.append(th)
    for r_ in t.rows:
        trPr = r_._tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); cs.set(qn('w:val'), 'true'); trPr.append(cs)
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(3); sp.paragraph_format.line_spacing = 0.6
    return t

def caption(text):
    par = doc.add_paragraph(); runs(par, text, 8.5, color=RGBColor(0x40, 0x40, 0x40), italic=True); par.paragraph_format.space_after = Pt(2); par.paragraph_format.keep_with_next = True; return par

def figure(fn, width, cap):
    par = doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER; par.paragraph_format.space_after = Pt(1); par.paragraph_format.keep_with_next = True
    par.add_run().add_picture(os.path.join(FIG, fn), width=Inches(width))
    c = doc.add_paragraph(); runs(c, cap, 8.5, color=RGBColor(0x40, 0x40, 0x40), italic=True); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after = Pt(6)

def callout(title, lines, fill='EAF1F8', size=9):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; borders(t, '9DB7D1', '6'); cell_margins(t, 60, 60, 110, 110)
    c = t.rows[0].cells[0]; shade(c, fill); c.width = Inches(6.45)
    p0 = c.paragraphs[0]; p0.paragraph_format.space_after = Pt(1); runs(p0, f'**{title}**', size, color=NAVY)
    for ln in lines:
        q = c.add_paragraph(); q.paragraph_format.space_after = Pt(1); q.paragraph_format.line_spacing = 1.0; runs(q, ln, size)
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(2); sp.paragraph_format.line_spacing = 0.5

def page_break(): doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---------- footer with page number
def footer():
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run('GoodBelly case study  |  Decision Sciences II, IIM Bangalore  |  page '); r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
    for typ, txt in (('begin', None), (None, 'PAGE'), ('separate', None), ('end', None)):
        r2 = fp.add_run(); r2.font.size = Pt(8); r2.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
        if typ:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), typ); r2._r.append(e)
        else:
            e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt; r2._r.append(e)
        if typ == 'separate': r2.add_text('1')
footer()
sec.different_first_page_header_footer = True

# =================================================================== COVER
for _ in range(3): doc.add_paragraph()
P('Indian Institute of Management Bangalore', 15, 0, 'c', color=NAVY).runs[0].bold = True
P('PGP 2026-28  |  Decision Sciences II', 11, 30, 'c')
t = P('GoodBelly: Do In-Store Promotions Drive Sales?', 25, 6, 'c', color=NAVY); t.runs[0].bold = True
P('A JASP regression analysis of sales reps, endcap displays and in-store demos across 126 Whole Foods stores', 12.5, 28, 'c', italic=True)
P('Case Study Report', 17, 4, 'c').runs[0].bold = True
P('Group # [____]', 13, 24, 'c')
tb = table(['Team member', 'Roll number'], [[f'[Name {i}]', '[Roll no.]'] for i in range(1, 6)], [3.6, 2.2], ['l', 'l'], size=10)
P('Submitted: October 6, 2026  |  Course: Decision Sciences II  |  Software: JASP', 10, 0, 'c')
page_break()

# =================================================================== EXECUTIVE SUMMARY
H('Executive Summary', 1)
P(f"**Problem.** GoodBelly must justify its in-store marketing spend. We analyse {n0(R['n'])} store-week observations ({R['stores']} Whole Foods stores, {R['weeks']} weekly observations, May 4 to July 13, 2010) to test whether regional sales reps, endcap displays and in-store demos raise weekly unit sales, and for how long.")
P('**Approach (all in JASP).** We followed a deliberate sequence, so that each model is justified by the failure of the one before it: '
  f"(1) a **base model** of units on all eight variables (R² = {n3(m1['r2'])}); (2) assumption checks, which it **fails** (Durbin-Watson {n2(m1['dw'])}, residual kurtosis {n1(m1['kurt'])}, a block of large negative residuals); "
  f"(3) a **log-level model** on ln(units) (R² = {n3(ml['r2'])}); (4) assumption checks again, which it **still fails** (kurtosis {n1(ml['kurt'])}, skew {n2(ml['skew'])}, unequal spread); "
  "(5) a screen of **16 interaction terms**, which singled out **Endcap × Sales Rep** (it lifts R² by 13 points; the next best adds under 1); "
  f"(6) **stepwise regression** on the 8 variables plus the 16 interactions, which gives our **final model** (R² = {n3(m2['r2'])}, all assumptions met); (7) answers to the six case questions.")
P('**Final model.** ' + f"Units = {n1(m2['coef']['Intercept']['b'])} + {n1(m2['coef']['Rep_x_Endcap']['b'])}·(Endcap × Rep) + {n1(m2['coef']['Rep']['b'])}·Rep + {n1(m2['coef']['Demo']['b'])}·Demo + {n1(m2['coef']['D13']['b'])}·Demo1-3 + {n1(m2['coef']['D45']['b'])}·Demo4-5 − {n1(-m2['coef']['Price']['b'])}·Price", after=3)
P('**Key findings**', after=2)
B(f"**Endcaps work only together with a regional rep.** In stores with a rep, an endcap adds about {n0(m2['coef']['Rep_x_Endcap']['b'])} units a week (95% CI {n0(m2['coef']['Rep_x_Endcap']['lo'])} to {n0(m2['coef']['Rep_x_Endcap']['hi'])}); in stores without one the effect is zero (p = 0.96). Average weekly units: {n0(R['groups']['11']['mean'])} (endcap with rep) vs {n0(R['groups']['10']['mean'])} (endcap without rep).")
B(f"**Demos work and the lift does not fade.** A demo adds {n0(m2['coef']['Demo']['b'])} units in the demo week and about {n0(m2['coef']['D13']['b'])} a week in weeks 1-3 and {n0(m2['coef']['D45']['b'])} in weeks 4-5+ (all p < 0.001); the two later periods are statistically identical (p = {n2(R['q3']['D13_eq_D45']['p'])}). One demo is worth about {n0(R['q3']['total'])} extra units over six weeks.")
B(f"**Regional reps help on their own** (+{n0(m2['coef']['Rep']['b'])} units a week) and are the precondition for endcaps. **Price** matters but demand is inelastic (about {n2(R['elasticity'])} at the mean). Store neighbourhood (competitors, fitness centres, region) adds nothing.")
P('**Recommendations.** (1) Fund endcaps only in rep-covered stores, and extend rep coverage first. (2) Keep the demo programme and judge it over a six-week window, not on demo-day sales. (3) Do not use price cuts to buy volume. (4) Test endcaps in a few more stores without reps: that finding rests on three stores. (5) Collect store traffic, demo cost and a longer time window to learn when the demo lift ends and to compute return on investment.', after=2)
page_break()

# =================================================================== 1 PROBLEM AND DATA
H('Background: problem and data', 1)
P(f"GoodBelly (a probiotic juice) spends its limited marketing budget on three programmes: **in-store demos** (a representative hands out samples for three hours), **endcap displays** (the end-of-aisle hub, promoted through rep and store competitions) and **regional sales reps** (face-to-face contact with the store, versus the national rep only). Management asks whether demos and endcaps actually raise sales and whether any demo boost is only temporary. The dataset has {n0(R['n'])} store-week rows. Weekly units average {n0(R['mean_units'])} (SD {n0(R['desc']['Units']['sd'])}, range {n0(R['desc']['Units']['min'])} to {n0(R['desc']['Units']['max'])}).")
table(['Variable', 'Type', 'Meaning', 'Summary'],
      [['Units (Y)', 'scale', 'weekly units sold per store', f"mean {n0(R['desc']['Units']['mean'])}, SD {n0(R['desc']['Units']['sd'])}"],
       ['Price', 'scale', 'average retail price ($)', f"mean {n2(R['desc']['Price']['mean'])}, range {n2(R['desc']['Price']['min'])} to {n2(R['desc']['Price']['max'])}"],
       ['Rep', 'dummy', '1 = regional sales rep in store', f"{n0(R['dummy_counts']['Rep'])} store-weeks ({R['dummy_counts']['Rep'] / R['n'] * 100:.0f}%)"],
       ['Endcap', 'dummy', '1 = endcap promotion', f"{R['dummy_counts']['Endcap']} store-weeks, {R['n_endcap_stores']} stores"],
       ['Demo', 'dummy', '1 = demo this week', f"{R['dummy_counts']['Demo']} store-weeks"],
       ['D13 (Demo1-3)', 'dummy', '1 = demo 1-3 weeks ago', f"{R['dummy_counts']['D13']} store-weeks"],
       ['D45 (Demo4-5)', 'dummy', '1 = demo 4-5 (or more) weeks ago', f"{R['dummy_counts']['D45']} store-weeks"],
       ['Natural, Fitness', 'count', 'natural retailers / fitness centres within 5 miles', f"means {n1(R['desc']['Natural']['mean'])} and {n1(R['desc']['Fitness']['mean'])}"]],
      [1.25, 0.6, 2.75, 1.85], ['l', 'l', 'l', 'l'])
P(f"Two features of the data drive everything that follows. First, endcaps are rare: only {R['dummy_counts']['Endcap']} of {n0(R['n'])} store-weeks ({R['dummy_counts']['Endcap'] / R['n'] * 100:.1f}%). Second, the three demo flags are a timeline for the same demo (this week / 1-3 weeks ago / 4-5+ weeks ago), so 'no recent demo' is the base group and each demo coefficient is the lift compared with no recent demo. A few stores ran repeat demos ({R['demo_overlap_rows']} store-weeks carry two flags), and their effects simply add.", after=3)
callout('Setup in JASP', ["Open **GoodBelly_JASP.csv** (1,386 rows; it holds the original columns plus the 16 interaction columns, D13_plus_D45 and lnUnits). Regression → Classical → **Linear Regression**. Every table in this report comes from that one analysis screen; Statistics and Plots options used are listed in Appendix A."])

# =================================================================== 2 BASE MODEL
H('Step 1: Base model', 1)
P('The natural first model puts all eight variables in, with no transformation and no interactions:')
P('Units = β₀ + β₁ Price + β₂ Rep + β₃ Endcap + β₄ Demo + β₅ D13 + β₆ D45 + β₇ Natural + β₈ Fitness + ε', 10, 4, 'c', italic=True)
rows = []
for v, lab in [('Intercept', 'Intercept'), ('Price', 'Price'), ('Rep', 'Rep'), ('Endcap', 'Endcap'), ('Demo', 'Demo'), ('D13', 'D13 (Demo 1-3)'), ('D45', 'D45 (Demo 4-5)'), ('Natural', 'Natural'), ('Fitness', 'Fitness')]:
    c = m1['coef'][v]; rows.append([lab, n2(c['b']), n2(c['se']), '' if v == 'Intercept' else n3(m1['std_beta'][v]), n2(c['t']), p_(c['p']), f"{n1(c['lo'])} to {n1(c['hi'])}", '' if v == 'Intercept' else n2(m1['vif'][v])])
caption('Table 1. Base model coefficients (JASP: Linear Regression, Method = Enter)')
table(['Variable', 'Coef.', 'Std. error', 'Std. β', 't', 'p', '95% CI', 'VIF'], rows, [1.25, 0.7, 0.75, 0.6, 0.6, 0.6, 1.35, 0.5], shade_rows={7: 'F3F3F3', 8: 'F3F3F3'})
table(['Model fit', 'R²', 'Adj. R²', 'RMSE (units)', 'F (8, 1377)', 'p', 'SS regression', 'SS residual'],
      [['Base model', n3(m1['r2']), n3(m1['ar2']), n1(m1['rmse']), n1(m1['F']), '<0.001', n0(m1['ss_reg']), n0(m1['ss_res'])]], [1.05, 0.55, 0.65, 0.85, 0.85, 0.5, 1.0, 1.0])
P('**What we observe.**', after=2, keep=True)
B(f"The model is useful overall (F = {n1(m1['F'])}, p < 0.001) and explains {m1['r2'] * 100:.0f}% of the variation in weekly units. Six of eight predictors are significant and have the expected signs: a higher price lowers sales; reps, endcaps and every demo period raise them.")
B(f"Endcap has by far the largest effect (+{n0(m1['coef']['Endcap']['b'])} units, standardised β = {n2(m1['std_beta']['Endcap'])}), followed by Rep (+{n0(m1['coef']['Rep']['b'])}) and the demo flags (+{n0(m1['coef']['D13']['b'])} to +{n0(m1['coef']['Demo']['b'])}). Holding the others fixed, a demo week adds about {n0(m1['coef']['Demo']['b'])} units and the weeks after still add about {n0(m1['coef']['D13']['b'])} and {n0(m1['coef']['D45']['b'])}.")
B(f"Natural retailers (p = {n2(m1['coef']['Natural']['p'])}) and Fitness centres (p = {n2(m1['coef']['Fitness']['p'])}) are not significant.")
B(f"Multicollinearity is not a problem: all VIFs are between {n2(min(m1['vif'].values()))} and {n2(max(m1['vif'].values()))} (rule of thumb: worry above 4).")
P('A decent R² and significant p-values are only trustworthy if the regression assumptions hold, so the next step is to test them.', after=2)

# =================================================================== 3 ASSUMPTIONS BASE
H('Step 2: Assumption checks on the base model', 1)
P('The p-values and confidence intervals in Table 1 assume that the errors have mean zero and constant variance, are independent, and are normally distributed, and that the predictors are not collinear. We check each with the JASP residual plots, Durbin-Watson statistic and casewise diagnostics.')
figure('diag_base.png', 5.55, 'Figure 1. Base model: residuals vs. predicted, histogram, normal P-P plot and Q-Q plot of the standardized residuals')
table(['Assumption', 'Check', 'Base model result', 'Verdict'],
      [['Linearity, constant variance', 'Residuals vs. predicted', f"uneven spread; a separate cluster of large negative residuals (down to {n0(m1['min_res'])}) at predicted values of 450-600", '**Fails**'],
       ['Independence', 'Durbin-Watson (about 2 is good)', f"{n2(m1['dw'])}: positive autocorrelation", '**Fails**'],
       ['Normality', 'Histogram, P-P, Q-Q, skewness, kurtosis', f"skew {n2(m1['skew'])}, kurtosis {n2(m1['kurt'])} (normal = 0 and 3); heavy left tail", '**Fails**'],
       ['No multicollinearity', 'VIF', f"all VIF ≤ {n2(max(m1['vif'].values()))}", 'Passes'],
       ['Outliers / influence', "Std. residual beyond ±3; Cook's D", f"{m1['n_sr3']} cases beyond ±3; max Cook's D {n3(m1['max_cook'])} (no influential point)", 'Warning']],
      [1.4, 1.55, 2.85, 0.65], ['l', 'l', 'l', 'c'])
P(f"**Reading the evidence.** The model fails three of the four core assumptions, even though its R² looks respectable. The problem is concentrated, not general: the residual plot shows one detached cluster far below zero, and {m1['sr3_in_endcap_norep']} of the {m1['n_sr3']} cases beyond ±3 standard deviations are endcap weeks (we return to these in Step 5). Because the errors are skewed and heavy-tailed, the base model's standard errors and p-values cannot be fully trusted. A common first remedy is to change the scale of the dependent variable.")

# =================================================================== 4 LOG LEVEL
H('Step 3: Log-level model', 1)
P('Weekly units are right-skewed (skewness {:.2f}) and the residuals of the base model are left-skewed and unevenly spread, which are the classic signs that a **log-level** model (ln Units on the same eight predictors in levels) might behave better. A coefficient b now means that units change by about 100·(e^b − 1)% when the variable rises by one unit.'.format(R['units_skew']))
rows = []
for v, lab in [('Intercept', 'Intercept'), ('Price', 'Price'), ('Rep', 'Rep'), ('Endcap', 'Endcap'), ('Demo', 'Demo'), ('D13', 'D13 (Demo 1-3)'), ('D45', 'D45 (Demo 4-5)'), ('Natural', 'Natural'), ('Fitness', 'Fitness')]:
    c = ml['coef'][v]; rows.append([lab, n3(c['b']), n3(c['se']), n2(c['t']), p_(c['p']), '' if v == 'Intercept' else f"{ml['pct'][v]:+.1f}%"])
caption('Table 2. Log-level model coefficients (dependent variable: ln Units)')
table(['Variable', 'Coef.', 'Std. error', 't', 'p', '% change in units'], rows, [1.5, 0.85, 0.85, 0.7, 0.7, 1.3], shade_rows={7: 'F3F3F3', 8: 'F3F3F3'})
P(f"The signs and significance are the same as before (Natural and Fitness again drop out). The fit is R² = {n3(ml['r2'])} (adjusted {n3(ml['ar2'])}); this is not directly comparable with the base model because the dependent variable has changed, but it is certainly not better. An endcap is now estimated at +{ml['pct']['Endcap']:.0f}% and a demo week at +{ml['pct']['Demo']:.0f}%. Before accepting this model, we test its assumptions.")

# =================================================================== 5 ASSUMPTIONS LOG
H('Step 4: Assumption checks on the log-level model', 1)
figure('diag_log.png', 5.55, 'Figure 2. Log-level model: residuals vs. predicted, histogram, normal P-P plot and Q-Q plot of the standardized residuals')
table(['Check', 'Base model', 'Log-level model', 'Target'],
      [['Durbin-Watson', n2(m1['dw']), n2(ml['dw']), 'about 2'],
       ['Skewness of residuals', n2(m1['skew']), n2(ml['skew']), 'about 0'],
       ['Kurtosis of residuals', n2(m1['kurt']), n2(ml['kurt']), 'about 3'],
       ['Cases beyond ±3 std. residuals', m1['n_sr3'], ml['n_sr3'], 'few'],
       ['Mean residual, endcap weeks without a rep (n = 17)', f"{s0(m1['res_by_group']['10'])} units", f"{ml['res_by_group']['10']:+.2f} (log units)", 'about 0'],
       ['Mean residual, endcap weeks with a rep (n = 36)', f"{s0(m1['res_by_group']['11'])} units", f"{ml['res_by_group']['11']:+.2f} (log units)", 'about 0']],
      [2.9, 1.1, 1.4, 1.0], ['l', 'r', 'r', 'r'])
P(f"**Verdict: the log transformation does not solve the problem.** It brings the Durbin-Watson statistic to {n2(ml['dw'])} and lowers the kurtosis from {n1(m1['kurt'])} to {n1(ml['kurt'])}, but the residuals remain clearly left-skewed ({n2(ml['skew'])}), still heavy-tailed, and the residual plot still shows the same detached group of points. The reason is visible in the last two rows: the same 17 endcap-without-rep weeks are still over-predicted (they sold about {(1 - 2.718281828 ** ml['res_by_group']['10']) * 100:.0f}% fewer units than the model expects) while the 36 endcap-with-rep weeks are under-predicted. A change of scale cannot fix a pattern that is **structural**: the base model treats an endcap as worth the same everywhere, and the data say it is not. This points to a missing interaction variable.")

# =================================================================== 6 INTERACTIONS
H('Step 5: Looking for interaction variables', 1)
P('An interaction term (the product of two variables) lets the effect of one variable depend on the level of another (Session 6). With eight predictors there are 28 possible pairwise products. We did not throw all 28 into the model: with that many candidates, some would look significant by chance, and many have no business meaning. We selected **16** on business logic and left out **12**.')
H('5a. Why these 16 interactions', 2)
table(['Group', 'Interactions (count)', 'Question each one asks'],
      [['A. Rep × promotion', 'Endcap × Rep, Rep × Demo, Rep × D13, Rep × D45 (4)', 'Does face-to-face contact make the other promotions work better? Reps set up demos and win endcap placement (the case describes the rep endcap competition).'],
       ['B. Endcap × demo', 'Endcap × Demo, Endcap × D13, Endcap × D45 (3)', 'Do an end-of-aisle display and in-store sampling reinforce each other?'],
       ['C. Price × promotion', 'Price × Rep, Endcap, Demo, D13, D45 (5)', 'Is price sensitivity different when a store is being promoted? (Matters for the pricing recommendation.)'],
       ['D. Neighbourhood × lever', 'Natural × Rep, Natural × Endcap, Fitness × Rep, Fitness × Demo (4)', 'Does competition (natural retailers) or a health-conscious catchment (fitness centres) change how well a lever works? One plausible pair per lever.']],
      [1.35, 2.4, 2.7], ['l', 'l', 'l'])
P('**The 12 we excluded, and why:** (i) Demo × D13, Demo × D45 and D13 × D45 (3 terms): these are alternative timing flags of the same demo, so their product is non-zero only in repeat-demo weeks (5 to 18 rows) and has no meaning of its own. (ii) Natural × Fitness (1): two neighbourhood counts with no promotion lever involved. (iii) Price × Natural and Price × Fitness (2): not related to any of the case questions. (iv) Natural × Demo, Natural × D13, Natural × D45, Fitness × Endcap, Fitness × D13, Fitness × D45 (6): repeat the same neighbourhood-moderates-lever idea already covered by group D, and add candidates without adding a new question. The 16 chosen terms keep the stepwise candidate pool at a manageable 24 (8 variables + 16 interactions). As a safeguard, we later repeat the search over all 154 two-, three- and four-way products (Step 6d).', size=10)
H('5b. Screening: adding each interaction to the base model', 2)
P('To see which interactions carry information, we add each one separately to the base model (JASP: add the product column as a ninth covariate and read R² change and its t-test). Table 3 ranks them.', after=2)
rows = [[r['term'].replace('Rep_x_Endcap', 'Endcap × Rep').replace('_x_', ' × '), r['nonzero'], n2(r['b']), n2(r['t']), p_(r['p']), f"{r['dr2'] * 100:.2f}"] for r in R['screen']]
caption('Table 3. Each of the 16 interactions added one at a time to the base model (base R² = %.4f)' % m1['r2'])
table(['Interaction added', 'Non-zero rows', 'Coef.', 't', 'p', 'R² gain (points)'], rows, [1.9, 1.0, 0.9, 0.7, 0.8, 1.15], shade_rows={0: 'FBE9E7'}, size=8.3)
H('5c. Why Endcap × Sales Rep is the one we choose', 2)
figure('why_endcap_rep.png', 5.7, 'Figure 3. Left: raw average units by endcap and rep. Right: how far off the base model is for each of those four groups')
B(f"**It dominates statistically.** It lifts R² from {n3(m1['r2'])} to {n3(R['screen'][0]['r2'])} (+{R['screen'][0]['dr2'] * 100:.1f} points, t = {n1(R['screen'][0]['t'])}). The runner-up adds {R['screen'][1]['dr2'] * 100:.1f} points. Every other term adds under 0.4.")
B(f"**It explains exactly what the base model got wrong.** Endcap weeks without a rep average only {n0(R['groups']['10']['mean'])} units (n = {R['groups']['10']['n']}), against {n0(R['groups']['11']['mean'])} with a rep (n = {R['groups']['11']['n']}). The base model predicts an endcap lifts sales everywhere, so it over-predicts the first group by {n0(-m1['res_by_group']['10'])} units on average and under-predicts the second by {n0(m1['res_by_group']['11'])}. Those weeks are the detached cluster in Figure 1 and account for {m1['sr3_in_endcap_norep']} of the {m1['n_sr3']} outliers.")
B(f"**It has a business reason.** An endcap has to be won, built and kept stocked by a person in the store. In the case, the endcap competition is run by the sales reps, who must persuade stores to place the product there. An endcap without face-to-face support is plausibly just a display.")
B(f"**The other 'significant' interactions look like proxies for it.** Endcap × Demo and Endcap × D13 look significant in Table 3, but they rest on only {R['screen'][1]['nonzero']} and {R['screen'][2]['nonzero']} non-zero rows, all inside endcap weeks; Natural × Rep (p = 0.0001 alone) is also no longer significant once Endcap × Rep is in the model. With Endcap × Rep included, none of the 15 others is significant (smallest p = {n2(min(R['q5']['other15'].values()))}, Step 6d).")
P(f"**Caution.** The 'no rep' endcap evidence comes from only three stores ({', '.join(R['endcap_norep_stores'])}) and 17 store-weeks, so it is strong but should be confirmed with more endcap-without-rep stores.", size=10)

# =================================================================== 7 STEPWISE FINAL
H('Step 6: Stepwise regression and the final model', 1)
H('6a. Stepwise selection', 2)
P(f"Rather than choosing variables by eye, we let the data pick them. In JASP we set **Method = Stepwise** with the 24 candidates (8 variables + 16 interactions), **entry p = 0.05 and removal p = 0.10**. At each step the algorithm adds the candidate with the largest partial correlation if its partial F-test has p < 0.05, then removes any variable whose p rises above 0.10, until nothing changes. Natural, Fitness, plain Endcap and all other 15 interactions never enter, and nothing is ever removed.", after=3)
rows = [[x['step'], x['var'].replace('Rep_x_Endcap', 'Endcap × Rep').replace('_x_', ' × '), n1(x['fchg']), n3(x['r2']), n3(x['dr2']), n3(x['ar2']), n1(x['rmse'])] for x in R['stepwise']['path']]
caption('Table 4. Stepwise path (JASP Model Summary, models M1 to M6; Endcap × Rep is the column Rep_x_Endcap)')
table(['Step', 'Variable entered', 'F change', 'R²', 'R² change', 'Adj. R²', 'RMSE'], rows, [0.5, 1.9, 0.9, 0.7, 0.9, 0.8, 0.75], ['c', 'l', 'r', 'r', 'r', 'r', 'r'])
P(f"Endcap × Rep enters first and by itself explains {R['stepwise']['path'][0]['r2'] * 100:.0f}% of the variation. Notice also that plain Endcap does not appear: its effect is entirely inside the interaction. (If stepwise is run on the eight main effects alone, it keeps plain Endcap, drops only Natural and Fitness, and so inherits the base model's problems.) We also ran stepwise with the same 24 candidates on ln Units: it needs {R['stepwise_ln']['n_terms']} terms and still gives R² = {n2(R['stepwise_ln']['r2'])}, skew {n2(R['stepwise_ln']['skew'])} and kurtosis {n1(R['stepwise_ln']['kurt'])}. The log scale adds nothing once the interaction is available, so **the final model is the stepwise model in level units.**", size=10)
H('6b. Final model', 2)
labs = {'Intercept': 'Intercept', 'Rep_x_Endcap': 'Endcap × Rep', 'Rep': 'Rep', 'D13': 'D13 (Demo 1-3)', 'Demo': 'Demo', 'D45': 'D45 (Demo 4-5)', 'Price': 'Price'}
rows = []
for v in ['Intercept', 'Rep_x_Endcap', 'Rep', 'Demo', 'D13', 'D45', 'Price']:
    c = m2['coef'][v]; rows.append([labs[v], n2(c['b']), n2(c['se']), '' if v == 'Intercept' else n3(m2['std_beta'][v]), n2(c['t']), p_(c['p']), f"{n1(c['lo'])} to {n1(c['hi'])}", '' if v == 'Intercept' else n2(m2['vif'][v])])
caption('Table 5. Final stepwise model (dependent variable: Units)')
table(['Variable', 'Coef.', 'Std. error', 'Std. β', 't', 'p', '95% CI', 'VIF'], rows, [1.25, 0.7, 0.75, 0.6, 0.6, 0.6, 1.35, 0.5])
table(['Model fit', 'R²', 'Adj. R²', 'RMSE (units)', 'F (6, 1379)', 'p', 'Durbin-Watson'],
      [['Final model', n3(m2['r2']), n3(m2['ar2']), n1(m2['rmse']), n1(m2['F']), '<0.001', n2(m2['dw'])]], [1.1, 0.7, 0.8, 1.0, 1.0, 0.6, 1.1])
P(f"**Units = {n1(m2['coef']['Intercept']['b'])} + {n1(m2['coef']['Rep_x_Endcap']['b'])} (Endcap × Rep) + {n1(m2['coef']['Rep']['b'])} Rep + {n1(m2['coef']['Demo']['b'])} Demo + {n1(m2['coef']['D13']['b'])} D13 + {n1(m2['coef']['D45']['b'])} D45 − {n1(-m2['coef']['Price']['b'])} Price**", 10, 3, 'c')
P(f"Compared with the base model, R² rises from {n3(m1['r2'])} to {n3(m2['r2'])} and the typical prediction error (RMSE) falls from {n1(m1['rmse'])} to {n1(m2['rmse'])} units. In 10-fold cross-validation with whole stores held out, the out-of-sample RMSE falls from {n1(R['cv_rmse']['base'])} to {n1(R['cv_rmse']['final'])}, so the improvement is not over-fitting. Standardised betas rank the levers: Endcap × Rep ({n2(m2['std_beta']['Rep_x_Endcap'])}) > Rep ({n2(m2['std_beta']['Rep'])}) > D13 ({n2(m2['std_beta']['D13'])}) > Demo ({n2(m2['std_beta']['Demo'])}) > D45 ({n2(m2['std_beta']['D45'])}) > Price ({n2(m2['std_beta']['Price'])}).", size=10)
H('6c. Assumption checks on the final model', 2)
figure('diag_final.png', 5.55, 'Figure 4. Final model: residuals vs. predicted, histogram, normal P-P plot and Q-Q plot of the standardized residuals')
table(['Check', 'Base model', 'Log-level model', 'Final model', 'Target'],
      [['Residuals vs. predicted', 'uneven, detached cluster', 'uneven, detached cluster', 'even band around zero', 'no pattern'],
       ['Durbin-Watson', n2(m1['dw']), n2(ml['dw']), n2(m2['dw']), 'about 2'],
       ['Skewness / kurtosis', f"{n2(m1['skew'])} / {n1(m1['kurt'])}", f"{n2(ml['skew'])} / {n1(ml['kurt'])}", f"{n2(m2['skew'])} / {n2(m2['kurt'])}", '0 / 3'],
       ['Largest VIF', n2(max(m1['vif'].values())), n2(max(ml['vif'].values())), n2(max(m2['vif'].values())), 'below 4'],
       ['Cases beyond ±3 std. residuals', m1['n_sr3'], ml['n_sr3'], m2['n_sr3'], 'few'],
       ["Max Cook's D", n3(m1['max_cook']), n3(ml['max_cook']), n3(m2['max_cook']), 'below 1'],
       ['Mean residual, endcap without rep', s0(m1['res_by_group']['10']) + ' units', f"{ml['res_by_group']['10']:+.2f} log", s0(m2['res_by_group']['10']) + ' units', 'about 0']],
      [2.15, 1.1, 1.15, 1.2, 0.85], ['l', 'r', 'r', 'r', 'r'], size=8.3)
P(f"**All assumptions are now reasonably met.** The residual plot is an even band, the P-P and Q-Q plots follow the diagonal, skewness and kurtosis are essentially those of a normal distribution, and the Durbin-Watson statistic is {n2(m2['dw'])}. Only {m2['n_sr3']} of {n0(R['n'])} cases ({m2['n_sr3'] / R['n'] * 100:.2f}%) lie beyond ±3, in line with chance, and no point is influential (largest Cook's D {n3(m2['max_cook'])}). Of the 20 highest-leverage cases, {m2['top20_lev_rep_x_endcap']} are endcap-with-rep weeks (rare, so each carries weight), but none distorts the fit. A Breusch-Pagan test run outside JASP gives p = {n3(m2['bp_p'])}, no evidence of unequal variance (it was p < 0.001 for the base and log models).", size=10)
H('6d. Robustness', 2)
B(f"**Store clustering.** The same stores are observed for 11 weeks, so errors may be correlated within a store. Re-estimating with standard errors clustered by store leaves every coefficient significant at p < 0.001 (for example Endcap × Rep: SE {n1(m2['cluster']['Rep_x_Endcap']['se'])} vs {n1(m2['coef']['Rep_x_Endcap']['se'])}; Demo: {n1(m2['cluster']['Demo']['se'])} vs {n1(m2['coef']['Demo']['se'])}).")
B(f"**Did we miss a better combination?** Repeating the stepwise search over all 2-, 3- and 4-way products of the eight variables (154 terms, 108 with enough non-zero rows) selects the same six terms in every pool (Python script analysis.py, supplied). Adding each of the 15 other interactions to the final model gives p-values from {n2(min(R['q5']['other15'].values()))} to {n2(max(R['q5']['other15'].values()))}.")

# =================================================================== 8 ANSWERS
H('Step 7: Answers to the case questions', 1)
H('Question 1. Build a statistical model of sales and promotional efforts. What do you observe?', 2)
P(f"**Answer.** We built the model in stages: the base model (R² = {n3(m1['r2'])}), a log-level version (R² = {n3(ml['r2'])}), and a stepwise model with an Endcap × Rep interaction. The final model (Table 5, R² = {n3(m2['r2'])}) is: **Units = {n1(m2['coef']['Intercept']['b'])} + {n1(m2['coef']['Rep_x_Endcap']['b'])} (Endcap × Rep) + {n1(m2['coef']['Rep']['b'])} Rep + {n1(m2['coef']['Demo']['b'])} Demo + {n1(m2['coef']['D13']['b'])} D13 + {n1(m2['coef']['D45']['b'])} D45 − {n1(-m2['coef']['Price']['b'])} Price.** What we observe:")
B(f"Every promotional lever has a positive and highly significant effect (all p < 0.001): a regional rep (+{n0(m2['coef']['Rep']['b'])} units a week), a demo week (+{n0(m2['coef']['Demo']['b'])}), the weeks after a demo (+{n0(m2['coef']['D13']['b'])} and +{n0(m2['coef']['D45']['b'])}) and, in rep stores, an endcap (+{n0(m2['coef']['Rep_x_Endcap']['b'])}).")
B(f"The endcap is not a stand-alone lever. The base model's +{n0(m1['coef']['Endcap']['b'])} is just the average of about +{n0(m2['coef']['Rep_x_Endcap']['b'])} in the 36 rep-store weeks and about zero in the 17 weeks without a rep (36/53 × {n0(m2['coef']['Rep_x_Endcap']['b'])} = {n0(R['blend_check']['expected'])}).")
B(f"Price has a negative effect: each extra $1 removes about {n0(-m2['coef']['Price']['b'])} units a week. At the mean price and sales this is an elasticity of {n2(R['elasticity'])}, i.e. demand is inelastic over the observed range (${n2(R['price_range'][0])} to ${n2(R['price_range'][1])}).")
B('Neighbourhood variables (Natural, Fitness) and region have no explanatory power once the promotion variables are in.')
H('Question 2. Does the model satisfy the regression assumptions? What additional analysis is needed?', 2)
P(f"**Answer.** The first two models do not; the final model does (Steps 2, 4 and 6c). The base model fails on independence (Durbin-Watson {n2(m1['dw'])}), constant variance and normality (kurtosis {n1(m1['kurt'])}); the log-level model improves the statistics but remains non-normal (skew {n2(ml['skew'])}, kurtosis {n1(ml['kurt'])}) because the cause was a missing interaction, not the scale. The final model has Durbin-Watson {n2(m2['dw'])}, skewness {n2(m2['skew'])}, kurtosis {n2(m2['kurt'])}, VIF ≤ {n2(max(m2['vif'].values()))}, an even residual band and no influential point. The additional analysis that verified it: residual plots and P-P/Q-Q plots, Durbin-Watson, VIF, casewise diagnostics with Cook's distance, a screen of interaction terms, stepwise selection, store-clustered standard errors and store-level cross-validation. One limitation remains: the same stores are observed repeatedly, which only a store random-effects model would handle fully (Question 5).")
H('Question 3. Does the in-store demo programme boost sales? For how long does the lift last?', 2)
P(f"**Answer: yes, and the lift lasts at least five weeks with no sign of fading.**", after=2)
B(f"A demo adds **+{n0(m2['coef']['Demo']['b'])} units** in the demo week (95% CI {n0(m2['coef']['Demo']['lo'])} to {n0(m2['coef']['Demo']['hi'])}), **+{n0(m2['coef']['D13']['b'])}** a week in weeks 1-3 afterwards and **+{n0(m2['coef']['D45']['b'])}** a week in weeks 4-5+ (all p < 0.001).")
B(f"Partial F-test, H₀: D13 = D45: F = {n2(R['q3']['D13_eq_D45']['F'])}, p = {n2(R['q3']['D13_eq_D45']['p'])}. The difference is {n1(R['q3']['diff_D45_D13']['est'])} units (95% CI {n0(R['q3']['diff_D45_D13']['lo'])} to +{n0(R['q3']['diff_D45_D13']['hi'])}), so we cannot reject that the lift in weeks 4-5+ equals that in weeks 1-3: **no decay**. The demo week itself is higher than the weeks after (Demo = D13: F = {n1(R['q3']['Demo_eq_D13']['F'])}, p < 0.001).")
B(f"One demo is therefore worth about {n0(m2['coef']['Demo']['b'])} + 3 × {n0(m2['coef']['D13']['b'])} + 2 × {n0(m2['coef']['D45']['b'])} = **{n0(R['q3']['total'])} extra units over six weeks**; only {m2['coef']['Demo']['b'] / R['q3']['total'] * 100:.0f}% of that falls on the demo day itself, so judging demos by demo-day sales understates them by a factor of four.")
B('Limit: D45 means "4-5 weeks ago or longer", so we know the lift persists to at least week 5 but not when it ends. The data span only 11 weeks.')
figure('levers.png', 5.7, 'Figure 5. Left: effect of each lever in the final model with 95% confidence intervals. Right: demo lift by period; the weeks 1-3 and 4-5+ bars are statistically equal')
H('Question 4. Does placement of the product within the store affect sales?', 2)
P(f"**Answer: yes, but only when a regional rep is present.** The endcap is the only placement variable in the data.", after=2)
B(f"With a rep, an endcap adds **+{n0(m2['coef']['Rep_x_Endcap']['b'])} units a week** (95% CI {n0(m2['coef']['Rep_x_Endcap']['lo'])} to {n0(m2['coef']['Rep_x_Endcap']['hi'])}); together with the rep's own +{n0(m2['coef']['Rep']['b'])}, such a store sells about {n0(R['q4']['rep_and_endcap'])} units more than a store with neither. Average weekly units: {n0(R['groups']['11']['mean'])} for endcap with rep versus {n0(R['groups']['00']['mean'])} for neither.")
B(f"Without a rep the endcap has no detectable effect: adding plain Endcap back to the final model gives a coefficient of {n1(R['q4']['plain_endcap']['b'])} (p = {n2(R['q4']['plain_endcap']['p'])}). Average units in those weeks were {n0(R['groups']['10']['mean'])}, close to the {n0(R['groups']['00']['mean'])} of stores with no endcap and no rep.")
B(f"Placement does not interact with demos (Endcap × Demo p = {n2(R['q5']['other15']['Endcap_x_Demo'])}, Endcap × D13 p = {n2(R['q5']['other15']['Endcap_x_D13'])}) or with price (p = {n2(R['q5']['other15']['Price_x_Endcap'])}).")
B(f"Caution: the 'no rep' result is based on {R['groups']['10']['n']} store-weeks from {len(R['endcap_norep_stores'])} stores ({', '.join(R['endcap_norep_stores'])}).")
H('Question 5. Are there suggestions to improve and refine the model?', 2)
P('**Answer.** Within the data, we tested the obvious refinements and none improves the final model (partial F-tests, each added to the final model):', after=2, keep=True)
table(['Refinement tested', 'Partial F', 'p', 'Adj. R²', 'Decision'],
      [['Add Natural + Fitness', n2(R['q5']['natural_fitness']['F']), n2(R['q5']['natural_fitness']['p']), n3(R['q5']['natural_fitness']['ar2']), 'not needed'],
       ['Add region dummies', n2(R['q5']['region']['F']), n2(R['q5']['region']['p']), n3(R['q5']['region']['ar2']), 'not needed'],
       ['Add Price²', n2(R['q5']['price_sq']['F']), n2(R['q5']['price_sq']['p']), n3(R['q5']['price_sq']['ar2']), 'price is linear'],
       ['Add a week trend', n2(R['q5']['week']['F']), n2(R['q5']['week']['p']), n3(R['q5']['week']['ar2']), 'no trend'],
       ['Add store dummies (fixed effects)', n2(R['q5']['store']['F']), n2(R['q5']['store']['p']), n3(R['q5']['store']['ar2']), 'no store effect left'],
       ['Add plain Endcap', '', n2(R['q4']['plain_endcap']['p']), '', 'inside the interaction'],
       ['Add any of the other 15 interactions', '', f"{n2(min(R['q5']['other15'].values()))} to {n2(max(R['q5']['other15'].values()))}", '', 'none significant']],
      [2.4, 0.8, 1.15, 0.8, 1.4], ['l', 'r', 'r', 'r', 'l'], size=8.5)
P('Improvements that need **new data or methods**: (i) store traffic and size, the most likely omitted variable (the case says demos work only in stores with enough foot traffic); (ii) competitor prices and promotions, holidays and seasonality (only 11 weeks, May to July); (iii) more endcap weeks in non-rep stores to confirm the interaction; (iv) a longer follow-up after each demo, with weekly dummies, to find when the lift ends; (v) a mixed (random-effects) model with a store effect to handle repeated observations; (vi) demo, rep and endcap costs, so that return on investment can be computed; (vii) a randomised pilot, because the data are observational and stores with reps and endcaps may differ in ways we cannot see.', size=10)
H('Question 6. Managerial implications and recommendations', 2)
table(['Finding', 'Recommendation'],
      [[f"Endcap adds about +{n0(m2['coef']['Rep_x_Endcap']['b'])} units a week, only in stores with a regional rep", '**Pair every endcap with a regional rep.** Extend rep coverage to endcap-ready stores first, and do not pay for endcaps (or the endcap competition prize) in stores without rep support.'],
       [f"A demo is worth about {n0(R['q3']['total'])} units over six weeks, with no fade through week 5", '**Keep the demo programme** and evaluate it over a six-week window, not demo-day sales. Compare the cost of a three-hour demo with 476 units × margin per unit to confirm the return on investment (cost data was not available).'],
       [f"A regional rep adds +{n0(m2['coef']['Rep']['b'])} units a week by itself", 'Add reps to uncovered stores; the rep both lifts sales directly and makes endcaps pay off.'],
       [f"Price elasticity about {n2(R['elasticity'])}", 'Do not use price cuts to buy volume: units rise less than proportionally, so revenue falls. Test modest price increases carefully (observed range is narrow).'],
       ['Natural, Fitness and region do not matter', 'Target stores by traffic and rep coverage, not by neighbourhood type.']],
      [2.55, 3.9], ['l', 'l'], size=8.8)
H('Limitations', 2)
P('The data are observational (stores were not randomly assigned to reps, endcaps or demos), cover one retailer and 11 weeks, and contain few endcap weeks, especially without a rep. The demo-lift duration is censored at "4-5+ weeks". No cost data are available, so we report incremental units, not profit. Results should be read as strong associations, to be confirmed with a pilot.', size=10)

# =================================================================== APPENDIX
H('Appendix A. JASP settings used at each step', 1)
table(['Step', 'JASP analysis and settings', 'Output used'],
      [['1 Base model', 'Regression → Classical → Linear Regression. Dependent: Units. Covariates: Price, Rep, Endcap, Demo, D13, D45, Natural, Fitness. Method: Enter. Statistics: Estimates, 95% CI, Model fit, R squared change, Collinearity (Tolerance/VIF), Durbin-Watson.', 'Table 1'],
       ['2 Assumptions', 'Same analysis. Residuals: Statistics, Casewise diagnostics (std. residual > 3, Cook\'s distance). Plots: Residuals vs. predicted, Residuals histogram (standardized), Q-Q plot standardized residuals.', 'Figure 1, assumption table'],
       ['3 Log-level', 'New Linear Regression; Dependent: lnUnits (column in the CSV, or compute ln(Units)); same eight covariates; same options.', 'Table 2'],
       ['4 Assumptions', 'Same plots and statistics as step 2.', 'Figure 2'],
       ['5 Interactions', 'Eight covariates plus one interaction column at a time (e.g. Rep_x_Endcap); read R squared change and the t-test of the added term.', 'Table 3'],
       ['6 Stepwise', 'New Linear Regression; Dependent: Units; Covariates: the 8 variables and 16 interaction columns; Method: Stepwise; Method Specification: Use p value, Entry 0.05, Removal 0.10; same Statistics and Plots.', 'Tables 4, 5; Figure 4'],
       ['Q3 test', 'Dependent: Units; covariates Rep_x_Endcap, Rep, Demo, D13_plus_D45, Price (D13 and D45 forced equal); compare its sum of squares with the final model: F = (SSE_reduced − SSE_full) / (SSE_full / 1379).', 'Q3 partial F']],
      [1.0, 4.35, 1.1], ['l', 'l', 'l'], size=8.3)
P('The JASP residual plots include the Q-Q plot; figures here were redrawn from the same residuals for a uniform layout and to add the P-P plot requested in the brief. The Breusch-Pagan test, store-clustered standard errors, cross-validation and the 154-term exhaustive search are supplementary checks run outside JASP (scripts jasp_report/analysis_jasp.py and analysis.py).', 9)

out = os.path.join(HERE, 'GoodBelly_JASP_Final_Report.docx')
doc.save(out); print('saved', out)
