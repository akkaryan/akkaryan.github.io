"""Builds the final submission report (docx) from results_jasp.json and fig/*.png.
Run analysis_jasp.py first, then this script, then convert to PDF with LibreOffice."""
import os, json
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__)); FIG = os.path.join(HERE, 'fig')
R = json.load(open(os.path.join(HERE, 'results_jasp.json')))
m1, ml, m2 = R['m1'], R['ml'], R['m2']
C1, C2, C3 = m1['coef'], ml['coef'], m2['coef']
BURG = RGBColor(0x7A, 0x1E, 0x2E); CHAR = RGBColor(0x33, 0x33, 0x33); GREYT = RGBColor(0x55, 0x55, 0x55)
HDR = '6B1B28'; BAND = 'F5F2EF'; FONT = 'Times New Roman'

def p_(p): return '<0.001' if p < 0.001 else f'{p:.3f}'
def n0(x): return f'{x:,.0f}'
def s0(x): return f'{x:+,.0f}'
def n1(x): return f'{x:,.1f}'
def n2(x): return f'{x:,.2f}'
def n3(x): return f'{x:.3f}'
MINUS = '−'
def neg(x, f=n1): return (MINUS + f(-x)) if x < 0 else f(x)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.left_margin = sec.right_margin = Inches(1.0); sec.top_margin = Inches(0.95); sec.bottom_margin = Inches(0.9)
sec.different_first_page_header_footer = True
st = doc.styles['Normal']; st.font.name = FONT; st.font.size = Pt(11); st.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'): rf.set(qn(a), FONT)
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.08; st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for name, size, color, before, after, italic in [('Heading 1', 15, BURG, 18, 6, False), ('Heading 2', 12, CHAR, 11, 3, False), ('Heading 3', 11, CHAR, 8, 2, True)]:
    h = doc.styles[name]; h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True; h.font.italic = italic; h.font.color.rgb = color
    hr = h.element.get_or_add_rPr(); f_ = hr.find(qn('w:rFonts'))
    if f_ is None: f_ = OxmlElement('w:rFonts'); hr.append(f_)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'): f_.set(qn(a), FONT)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if f_.get(qn(a)) is not None: del f_.attrib[qn(a)]
    h.paragraph_format.space_before = Pt(before); h.paragraph_format.space_after = Pt(after); h.paragraph_format.keep_with_next = True
    h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

def rule(par, color='7A1E2E', sz='6'):
    pPr = par._p.get_or_add_pPr(); b = OxmlElement('w:pBdr'); x = OxmlElement('w:bottom')
    x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), sz); x.set(qn('w:space'), '2'); x.set(qn('w:color'), color); b.append(x); pPr.append(b)

def runs(par, text, size=None, color=None, italic=None, bold=None):
    for i, seg in enumerate(text.split('**')):
        if not seg: continue
        r = par.add_run(seg); r.bold = (i % 2 == 1) or bool(bold)
        if italic: r.italic = True
        if size: r.font.size = Pt(size)
        if color is not None: r.font.color.rgb = color
    return par

def P(text, size=None, after=None, align=None, italic=None, keep=False, color=None, before=None):
    par = doc.add_paragraph(); runs(par, text, size, color, italic)
    if after is not None: par.paragraph_format.space_after = Pt(after)
    if before is not None: par.paragraph_format.space_before = Pt(before)
    if align == 'c': par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if align == 'l': par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if keep: par.paragraph_format.keep_with_next = True
    return par

def bullets(items, style='List Bullet', size=None):
    for t in items:
        par = doc.add_paragraph(style=style); runs(par, t, size); par.paragraph_format.space_after = Pt(3)
        par.paragraph_format.left_indent = Inches(0.3); par.paragraph_format.first_line_indent = Inches(-0.22)

def H1(t): h = doc.add_heading(t, 1); rule(h); return h
def H2(t): return doc.add_heading(t, 2)

def tcell_borders(tbl):
    tblPr = tbl._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for e, (v, sz, col) in dict(top=('single', '10', HDR), bottom=('single', '10', HDR), insideH=('single', '4', 'C8C2BC'), left=('nil', '0', 'auto'), right=('nil', '0', 'auto'), insideV=('nil', '0', 'auto')).items():
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'), v); x.set(qn('w:sz'), sz); x.set(qn('w:space'), '0'); x.set(qn('w:color'), col); b.append(x)
    anchor = next((x for x in (tblPr.find(qn('w:tblLayout')), tblPr.find(qn('w:tblCellMar')), tblPr.find(qn('w:tblLook'))) if x is not None), None)
    anchor.addprevious(b) if anchor is not None else tblPr.append(b)
    m = OxmlElement('w:tblCellMar')
    for k, v in (('top', 35), ('left', 80), ('bottom', 35), ('right', 80)):
        x = OxmlElement(f'w:{k}'); x.set(qn('w:w'), str(v)); x.set(qn('w:type'), 'dxa'); m.append(x)
    look = tblPr.find(qn('w:tblLook')); look.addprevious(m) if look is not None else tblPr.append(m)

def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hex_); tcPr.append(sh)

AL = {'l': WD_ALIGN_PARAGRAPH.LEFT, 'r': WD_ALIGN_PARAGRAPH.RIGHT, 'c': WD_ALIGN_PARAGRAPH.CENTER}
def caption(text):
    par = doc.add_paragraph(); runs(par, text, 10, bold=True, color=CHAR); par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    par.paragraph_format.space_before = Pt(6); par.paragraph_format.space_after = Pt(3); par.paragraph_format.keep_with_next = True

def table(header, rows, widths, align=None, size=9.5, bold_rows=(), shade_rows=None, cap=None, band=True):
    if cap: caption(cap)
    t = doc.add_table(rows=1 + len(rows), cols=len(header)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False; tcell_borders(t)
    align = align or ['l'] + ['r'] * (len(header) - 1)
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]; shade(c, HDR); c.width = Inches(widths[j]); par = c.paragraphs[0]; par.paragraph_format.space_after = Pt(0); par.paragraph_format.line_spacing = 1.0
        r = par.add_run(h); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(255, 255, 255); par.alignment = AL[align[j]]
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.rows[i + 1].cells[j]; c.width = Inches(widths[j]); par = c.paragraphs[0]; par.paragraph_format.space_after = Pt(0); par.paragraph_format.line_spacing = 1.0
            runs(par, str(v), size, bold=(i in bold_rows)); par.alignment = AL[align[j]]
            if shade_rows and i in shade_rows: shade(c, shade_rows[i])
            elif band and i % 2 == 1: shade(c, BAND)
    for j, w_ in enumerate(widths): t.columns[j].width = Inches(w_)
    for r_ in t.rows[:-1]:
        for c_ in r_.cells:
            for q_ in c_.paragraphs: q_.paragraph_format.keep_with_next = True
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); trPr.append(th)
    for r_ in t.rows:
        trPr = r_._tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); cs.set(qn('w:val'), 'true'); trPr.append(cs)
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(2); sp.paragraph_format.line_spacing = 0.5
    return t

def figure(fn, width, cap):
    par = doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER; par.paragraph_format.space_after = Pt(1); par.paragraph_format.keep_with_next = True
    par.add_run().add_picture(os.path.join(FIG, fn), width=Inches(width))
    c = doc.add_paragraph(); runs(c, cap, 10, bold=True, color=CHAR); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after = Pt(8)

def page_break(): doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def footer():
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run('GoodBelly Case Report  |  Section E, Group 15  |  '); r.font.size = Pt(9); r.font.color.rgb = GREYT
    for typ, txt in (('begin', None), (None, 'PAGE'), ('separate', None), ('end', None)):
        r2 = fp.add_run(); r2.font.size = Pt(9); r2.font.color.rgb = GREYT
        if typ:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), typ); r2._r.append(e)
        else:
            e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt; r2._r.append(e)
        if typ == 'separate': r2.add_text('1')
footer()

# ================================================================= COVER
for _ in range(2): doc.add_paragraph()
lg = doc.add_paragraph(); lg.alignment = WD_ALIGN_PARAGRAPH.CENTER; lg.add_run().add_picture(os.path.join(HERE, 'assets', 'iimb_logo.png'), width=Inches(1.55))
P('Indian Institute of Management Bangalore', 13, 20, 'c', color=GREYT, before=6).runs[0].bold = True
t = P('GoodBelly: Case Report', 27, 4, 'c', color=BURG); t.runs[0].bold = True
sub = P('Using Statistics to Justify the Marketing Expense', 14, 26, 'c', color=CHAR); sub.runs[0].italic = True
P('Decision Sciences II  |  PGP 2026–28', 12, 3, 'c', color=CHAR)
P('Submitted to: Prof. Rajluxmi Murthy', 12, 22, 'c', color=CHAR)
P('Section E, Group 15', 13, 6, 'c', color=CHAR).runs[0].bold = True
table(['Team member', 'Roll number'], [['Harsh Jain', '2611356'], ['Anshh Chaturvedi', '2611355'], ['Snevi Kothari', '2611312'], ['Vishakha Tomar', '2611313'], ['Gokul Krishnan A', '2611367'], ['Yash Rajesh', '2611342']],
      [3.0, 1.8], ['l', 'l'], size=11, band=False)
P('6 October 2026', 11, 0, 'c', color=GREYT, before=14)
page_break()

# ================================================================= EXECUTIVE SUMMARY
H1('Executive Summary')
P(f"GoodBelly puts most of its small marketing budget into two in-store programmes at Whole Foods: product demonstrations and endcap displays. By July 2010 management was asking whether they lift sales, whether the lift lasts, and whether they are worth the money. We answer using {n0(R['n'])} weekly observations from {R['stores']} stores over {R['weeks']} weeks, analysed in JASP.")
P(f"We began with a regression of weekly units on all eight variables (Model 1). It explains {m1['r2'] * 100:.0f}% of the variation in sales and nearly every coefficient is significant, yet the residuals show clear problems: autocorrelation (Durbin-Watson {n2(m1['dw'])}), a heavy lower tail (kurtosis {n1(m1['kurt'])}) and a separate cluster of large negative residuals. Re-estimating the model on the log of units (Model 2) did not help; the same observations stayed badly fitted and the residuals remained non-normal. The trouble came from one place, the {R['dummy_counts']['Endcap']} store-weeks with an endcap. Endcap stores with a regional sales rep sold far more than predicted, endcap stores without one sold far less.")
P(f"We screened 16 plausible interaction terms and found that Endcap × Sales Rep outperformed all others by a wide margin. Stepwise regression on the original variables plus the 16 interactions then produced our final model (Model 3). It explains {m2['r2'] * 100:.1f}% of the variation, and its residuals meet the regression assumptions (Durbin-Watson {n2(m2['dw'])}, skewness {n2(m2['skew'])}, kurtosis {n2(m2['kurt'])}).")
P(f"Demos work, and the effect does not fade. A demo adds about {n0(C3['Demo']['b'])} units in the week it is held (roughly {C3['Demo']['b'] / R['baseline_no_demo_no_endcap'] * 100:.0f}% above a normal week) and a further {n0(C3['D13']['b'])} to {n0(C3['D45']['b'])} units a week for at least five weeks afterwards. Over six weeks that adds up to about {n0(R['q3']['total'])} extra units per demo, and more than three quarters of it arrives after the demo week. Endcaps are different. They add about {n0(C3['Rep_x_Endcap']['b'])} units a week, but only in stores with a regional sales rep; without a rep an endcap barely moves sales ({n0(R['groups']['10']['mean'])} units against {n0(R['groups']['00']['mean'])}). A regional rep is worth about {n0(C3['Rep']['b'])} units a week even without a promotion, and each $1 increase in price costs about {n0(-C3['Price']['b'])} units.")
P(f"Our advice is to keep the demo programme and to judge it over six weeks rather than one. Endcaps should go only into stores with a regional rep, and rep coverage deserves to be seen as the investment that makes endcaps pay. One number is still missing before the budget is settled, and that is cost: a demo pays for itself if it costs less than {n0(R['q3']['total'])} units times the margin per unit, so marketing needs to supply the cost of each programme and that margin.")
page_break()

# ================================================================= 1 INTRODUCTION
H1('1. Introduction')
H2('1.1 Case background')
P("GoodBelly, a probiotic juice brand from NextFoods Inc., cannot afford national advertising, so it relies on two in-store promotions at Whole Foods: demonstrations, in which trained representatives hand out samples, and endcap displays at the end of an aisle. Some stores also have a regional sales representative who visits in person, while the others are served only by the national representative. In July 2010 management began to question whether the promotions lift sales, whether any lift lasts, and whether they justify their cost. The marketing manager has to defend the budget with statistical evidence.")
H2('1.2 Data and variables')
P(f"The dataset has {n0(R['n'])} store-week observations from {R['stores']} Whole Foods stores between 4 May and 13 July 2010 ({R['weeks']} weekly dates). Weekly units average {n0(R['mean_units'])}, with a range of {n0(R['desc']['Units']['min'])} to {n0(R['desc']['Units']['max'])}. Only {R['dummy_counts']['Endcap']} store-weeks ({R['dummy_counts']['Endcap'] / R['n'] * 100:.1f}%) have an endcap, spread over {R['n_endcap_stores']} stores, a fact that matters a great deal later. In the JASP file the demo variables are named D13 and D45; we call them Demo1-3 and Demo4-5 in this report.", keep=True)
table(['Variable', 'Role', 'Definition'],
      [['Units (weekly sales)', 'Dependent', 'Units sold per store per week'],
       ['Demo', 'Promotional', '1 if the store had a demo in the current week'],
       ['Demo1-3', 'Promotional', '1 if the store had a demo 1 to 3 weeks ago'],
       ['Demo4-5', 'Promotional', '1 if the store had a demo 4 to 5 weeks ago (or earlier)'],
       ['Endcap', 'Promotional', '1 if the store took part in an endcap promotion'],
       ['Sales Rep', 'Control', '1 if the store has a regional sales rep (face-to-face contact), 0 if only the national rep'],
       ['Price', 'Control', f"Average retail price per store per week (${n2(R['price_range'][0])} to ${n2(R['price_range'][1])})"],
       ['Natural, Fitness', 'Control', 'Number of other natural retailers and of fitness centres within 5 miles']],
      [1.5, 1.0, 3.9], ['l', 'l', 'l'], cap='Table 1.1: Variables used in the analysis')
P("Demo, Demo1-3 and Demo4-5 describe how recently a store held a demo, so “no recent demo” is the base group and each of their coefficients is a lift relative to it. Repeat demos set more than one flag in " + f"{R['demo_overlap_rows']} store-weeks, and their effects simply add.")
H2('1.3 Approach')
P("We build the model in stages, so that each step is justified by the failure of the one before it: a base model with all variables, a check of its assumptions, a log-level model, a second check of the assumptions, a search for interaction terms, and finally a stepwise model that becomes our final model. The case questions are answered in Section 8, using that final model.")

# ================================================================= 2 MODEL 1
H1('2. Model 1: The Base Model')
P(f"The natural starting point is an ordinary least squares regression of weekly units on all eight variables, with no transformation and no interactions. The model is significant as a whole (F = {n1(m1['F'])}, p < 0.001) and explains {m1['r2'] * 100:.1f}% of the variation in weekly sales (adjusted R² = {n3(m1['ar2'])}, standard error {n1(m1['rmse'])} units).")
rows = []
for v, lab in [('Intercept', 'Intercept'), ('Price', 'Price'), ('Rep', 'Sales Rep'), ('Endcap', 'Endcap'), ('Demo', 'Demo'), ('D13', 'Demo1-3'), ('D45', 'Demo4-5'), ('Natural', 'Natural'), ('Fitness', 'Fitness')]:
    c = C1[v]; rows.append([lab, neg(c['b'], n2), n2(c['se']), '' if v == 'Intercept' else neg(m1['std_beta'][v], n3), neg(c['t'], n2), p_(c['p']), '' if v == 'Intercept' else n2(m1['vif'][v])])
table(['Variable', 'Coefficient', 'Std. error', 'Std. beta', 't', 'p-value', 'VIF'], rows, [1.35, 0.95, 0.85, 0.8, 0.7, 0.8, 0.6], cap='Table 2.1: Model 1 coefficients (dependent variable: Units)', shade_rows={7: 'FBF1B8', 8: 'FBF1B8'}, band=False)
table(['Source', 'Sum of squares', 'df', 'Mean square', 'F', 'p-value'],
      [['Regression', n0(m1['ss_reg']), m1['df_reg'], n0(m1['ss_reg'] / m1['df_reg']), n1(m1['F']), '<0.001'], ['Residual', n0(m1['ss_res']), f"{m1['df_res']:,}", n0(m1['ss_res'] / m1['df_res']), '', ''], ['Total', n0(m1['ss_tot']), f"{m1['df_reg'] + m1['df_res']:,}", '', '', '']],
      [1.3, 1.4, 0.7, 1.3, 0.8, 0.9], cap='Table 2.2: ANOVA for Model 1')
P(f"Six of the eight variables are significant at the 1% level and have the signs we would expect: a higher price lowers sales, while a regional rep, an endcap and each of the three demo periods raise them. The endcap has the largest effect (+{n0(C1['Endcap']['b'])} units, standardised beta {n2(m1['std_beta']['Endcap'])}), followed by the regional rep (+{n0(C1['Rep']['b'])}) and the demo variables (+{n0(C1['D45']['b'])} to +{n0(C1['Demo']['b'])}). The number of natural retailers (p = {n2(C1['Natural']['p'])}) and fitness centres (p = {n2(C1['Fitness']['p'])}) nearby are not significant, and their 95% confidence intervals include zero. Multicollinearity is not a concern: the variance inflation factors range from {n2(min(m1['vif'].values()))} to {n2(max(m1['vif'].values()))}, far below the usual warning level of 4 to 5.")
P("A good fit and significant coefficients are only worth trusting if the assumptions of linear regression hold. We check them next.")

# ================================================================= 3 ASSUMPTIONS M1
H1('3. Checking the Assumptions for Model 1')
P("The residual plots and the Durbin-Watson statistic from JASP let us assess linearity, constant variance, independence and normality. If the assumptions hold, the residuals should form a single random band of constant width around zero and follow the normal distribution.")
figure('diag_base.png', 5.6, 'Figure 3.1: Residual diagnostics for Model 1')
H2('3.1 Linearity and constant variance')
rn, rr, ro = m1['res_range_endcap_norep'], m1['res_range_endcap_rep'], m1['res_range_other']
P(f"Most of the residuals form a compact band around zero, but not all of them. A separate group of points sits far below the line, at predicted values of roughly 450 to 600 units, with residuals between {neg(rn[1], n0)} and {neg(rn[0], n0)}. Another group lies well above the line, with residuals between {s0(rr[0])} and {s0(rr[1])}. Both groups consist entirely of endcap store-weeks. The residuals of the other {R['n'] - R['dummy_counts']['Endcap']:,} observations stay within about {n0(max(abs(ro[0]), abs(ro[1])))} units of zero (standard deviation {n0(R['sd_by_endcap']['m1'][0])}), while those of the endcap weeks have a standard deviation of {n0(R['sd_by_endcap']['m1'][1])}, about four times larger. Linearity and constant variance both **fail**: Model 1 gives every endcap the same effect of +{n0(C1['Endcap']['b'])} units, which is too low for one group of stores and far too high for the other.")
H2('3.2 Independence')
P(f"The Durbin-Watson statistic is {n2(m1['dw'])}, well below 2, which points to positive autocorrelation. This is what the residual plot would lead us to expect. The same {R['stores']} stores are observed for {R['weeks']} weeks, and a store whose endcap is badly predicted in one week is badly predicted in the next. Independence **fails**.")
H2('3.3 Normality')
P(f"The histogram and the Q-Q plot show a long lower tail. Residuals reach {neg(m1['sr_range'][0], n1)} standard deviations, where a normal distribution would stop near −3, and the skewness and kurtosis are {neg(m1['skew'], n2)} and {n2(m1['kurt'])} (normal values are 0 and 3). {m1['n_sr3']} residuals lie beyond three standard deviations against the four or so expected in a sample this size, and every one of them is an endcap observation. Normality **fails**.")
table(['Assumption', 'Evidence from Model 1', 'Result'],
      [['Linearity', 'Endcap residuals split into two groups far from zero', 'Fails'], ['Constant variance', f"Residual SD {n0(R['sd_by_endcap']['m1'][0])} without endcap, {n0(R['sd_by_endcap']['m1'][1])} with", 'Fails'],
       ['Independence', f"Durbin-Watson {n2(m1['dw'])}", 'Fails'], ['Normality', f"Skewness {neg(m1['skew'], n2)}, kurtosis {n2(m1['kurt'])}, {m1['n_sr3']} residuals beyond ±3 SD", 'Fails'],
       ['No multicollinearity', f"Largest VIF {n2(max(m1['vif'].values()))}", 'Satisfied']],
      [1.6, 3.8, 1.0], ['l', 'l', 'l'], cap='Table 3.1: Assumption checks for Model 1')
P("Model 1 fits well by R² but fails four of the five checks, and each failure leads back to the same endcap observations. A common first remedy for skewed residuals is to change the scale of the dependent variable, so we try that next.")

# ================================================================= 4 MODEL 2 LOG
H1('4. Model 2: The Log-Level Model')
P(f"Weekly units are right-skewed (skewness {n2(R['units_skew'])}) and the residuals of Model 1 are left-skewed, which is the classic situation in which a log transformation of the dependent variable can help. Model 2 regresses the natural log of units on the same eight variables, which are left in their original levels. A coefficient b now means that units change by about 100·(e^b − 1)% when the variable increases by one unit.")
rows = []
for v, lab in [('Intercept', 'Intercept'), ('Price', 'Price'), ('Rep', 'Sales Rep'), ('Endcap', 'Endcap'), ('Demo', 'Demo'), ('D13', 'Demo1-3'), ('D45', 'Demo4-5'), ('Natural', 'Natural'), ('Fitness', 'Fitness')]:
    c = C2[v]; rows.append([lab, neg(c['b'], n3), n3(c['se']), neg(c['t'], n2), p_(c['p']), '' if v == 'Intercept' else f"{ml['pct'][v]:+.1f}%".replace('-', MINUS)])
table(['Variable', 'Coefficient', 'Std. error', 't', 'p-value', 'Change in units'], rows, [1.4, 1.0, 1.0, 0.9, 0.9, 1.2], cap='Table 4.1: Model 2 coefficients (dependent variable: ln Units)', shade_rows={7: 'FBF1B8', 8: 'FBF1B8'}, band=False)
P(f"The signs and the pattern of significance are the same as in Model 1, with Natural and Fitness again not significant. Model 2 explains {ml['r2'] * 100:.1f}% of the variation in ln Units (adjusted R² = {n3(ml['ar2'])}). That figure cannot be compared directly with Model 1, because the dependent variable has changed, but it certainly is not an improvement. The model puts the endcap effect at +{ml['pct']['Endcap']:.0f}% and the demo week at +{ml['pct']['Demo']:.0f}%. Before accepting it, we test its assumptions.")

# ================================================================= 5 ASSUMPTIONS M2
H1('5. Checking the Assumptions for Model 2')
figure('diag_log.png', 5.6, 'Figure 5.1: Residual diagnostics for Model 2')
P(f"The log transformation changes some statistics but not the picture. The Durbin-Watson statistic rises to {n2(ml['dw'])} and the kurtosis falls from {n1(m1['kurt'])} to {n1(ml['kurt'])}. However, the residuals are still clearly left-skewed ({neg(ml['skew'], n2)}), the lower tail of the Q-Q plot still bends away from the line (down to {neg(ml['sr_range'][0], n1)} standard deviations), {ml['n_sr3']} residuals lie beyond ±3, and the residual plot shows the same detached group of points. The reason is in the last two rows of Table 5.1. The 17 endcap weeks without a rep are still over-predicted (the stores sold about {(1 - 2.718281828 ** ml['res_by_group']['10']) * 100:.0f}% fewer units than the model expects), while the 36 endcap weeks with a rep are under-predicted.")
table(['Check', 'Model 1', 'Model 2 (log)', 'Target'],
      [['Durbin-Watson', n2(m1['dw']), n2(ml['dw']), 'about 2'], ['Skewness of residuals', neg(m1['skew'], n2), neg(ml['skew'], n2), 'about 0'], ['Kurtosis of residuals', n2(m1['kurt']), n2(ml['kurt']), 'about 3'],
       ['Residuals beyond ±3 SD', m1['n_sr3'], ml['n_sr3'], 'few'],
       ['Mean residual, endcap without rep (17 weeks)', neg(m1['res_by_group']['10'], n0) + ' units', neg(ml['res_by_group']['10'], n2) + ' (log units)', 'about 0'],
       ['Mean residual, endcap with rep (36 weeks)', s0(m1['res_by_group']['11']) + ' units', f"{ml['res_by_group']['11']:+.2f} (log units)", 'about 0']],
      [3.0, 1.0, 1.4, 1.0], ['l', 'r', 'r', 'r'], cap='Table 5.1: Model 1 and Model 2 compared')
P("A change of scale cannot repair a problem that is structural. Both models assume that an endcap is worth the same in every store, and the residuals say it is not. The log transformation was therefore the wrong remedy, and we turn to the likelier cause, a missing interaction between the endcap and something else about the store.")

# ================================================================= 6 INTERACTIONS
H1('6. Searching for Interaction Terms')
P("An interaction term is the product of two variables. It lets the effect of one variable depend on the level of the other. With eight variables there are 28 possible pairwise products. Putting all of them in would invite chance findings and include many pairs with no sensible meaning, so we chose 16 on business grounds and left out 12.")
H2('6.1 The 16 candidate interactions')
table(['Group', 'Interactions', 'Question asked'],
      [['A. Sales Rep with promotions (4)', 'Endcap × Sales Rep, Demo × Sales Rep, Demo1-3 × Sales Rep, Demo4-5 × Sales Rep', 'Does face-to-face contact make the other promotions work better? Reps set up demos and win endcap space; the case describes the rep endcap competition.'],
       ['B. Endcap with demos (3)', 'Endcap × Demo, Endcap × Demo1-3, Endcap × Demo4-5', 'Do an end-of-aisle display and in-store sampling reinforce each other?'],
       ['C. Price with promotions (5)', 'Price × Sales Rep, Endcap, Demo, Demo1-3, Demo4-5', 'Does price sensitivity change when a store is being promoted?'],
       ['D. Neighbourhood with levers (4)', 'Natural × Sales Rep, Natural × Endcap, Fitness × Sales Rep, Fitness × Demo', 'Do competition or a health-conscious catchment change how well a lever works? One plausible pair per lever.']],
      [1.55, 2.35, 2.5], ['l', 'l', 'l'], size=9, cap='Table 6.1: Interaction candidates and the reason for each group')
P("The 12 products we excluded are of three kinds. Demo × Demo1-3, Demo × Demo4-5 and Demo1-3 × Demo4-5 are different timing flags of the same demo, so their product is non-zero only in a handful of repeat-demo weeks (5 to 18 rows) and means nothing on its own. Natural × Fitness, Price × Natural and Price × Fitness combine neighbourhood counts with each other or with price and have no bearing on the case questions. The remaining six (Natural with Demo, Demo1-3 or Demo4-5, and Fitness with Endcap, Demo1-3 or Demo4-5) repeat the idea already covered by group D. Keeping the pool to 16 also keeps the stepwise candidate list at a manageable 24 terms.")
H2('6.2 Screening the candidates')
P("We added each interaction on its own to Model 1 and noted the t-test of the new term and the gain in R² (Table 6.2).", keep=True)
rows = [[r['term'].replace('Rep_x_Endcap', 'Endcap × Sales Rep').replace('_x_', ' × ').replace('Rep', 'Sales Rep').replace('D13', 'Demo1-3').replace('D45', 'Demo4-5').replace('Sales Sales Rep', 'Sales Rep'), r['nonzero'], neg(r['b'], n2), neg(r['t'], n2), p_(r['p']), f"{r['dr2'] * 100:.2f}"] for r in R['screen']]
table(['Interaction added to Model 1', 'Non-zero rows', 'Coefficient', 't', 'p-value', 'R² gain (points)'], rows, [2.3, 1.0, 0.95, 0.7, 0.8, 1.0], size=9, cap=f"Table 6.2: Each interaction added separately to Model 1 (Model 1 R² = {n3(m1['r2'])})", shade_rows={0: 'EBD3D6'}, band=False)
H2('6.3 Why Endcap × Sales Rep')
figure('why_endcap_rep.png', 5.7, 'Figure 6.1: Raw average units by endcap and sales rep (left) and the mean residual of Model 1 in each group (right)')
P(f"Endcap × Sales Rep is the clear choice. Adding it lifts R² from {n3(m1['r2'])} to {n3(R['screen'][0]['r2'])}, a gain of {R['screen'][0]['dr2'] * 100:.1f} points (t = {n1(R['screen'][0]['t'])}); the next best term adds {R['screen'][1]['dr2'] * 100:.1f} points. It also accounts for exactly what Model 1 got wrong. Endcap weeks without a rep average only {n0(R['groups']['10']['mean'])} units (n = {R['groups']['10']['n']}), compared with {n0(R['groups']['11']['mean'])} with a rep (n = {R['groups']['11']['n']}). Model 1 applies one endcap coefficient to both groups, so it over-predicts the first by {n0(-m1['res_by_group']['10'])} units on average and under-predicts the second by {n0(m1['res_by_group']['11'])}. These are the two detached clusters in Figure 3.1. The interaction also makes business sense: an endcap has to be won, built and kept stocked by someone in the store, and in the case it is the sales representatives who compete to place GoodBelly on endcaps.")
P(f"Two of the other terms look significant in Table 6.2, Endcap × Demo and Endcap × Demo1-3, but they rest on only {R['screen'][1]['nonzero']} and {R['screen'][2]['nonzero']} non-zero rows, all within endcap weeks, and they do not survive once Endcap × Sales Rep is in the model (p = {n2(R['q5']['other15']['Endcap_x_Demo'])} and {n2(R['q5']['other15']['Endcap_x_D13'])}). Natural × Sales Rep behaves the same way (p = {n2(R['q5']['other15']['Natural_x_Rep'])} after). The evidence for the “no rep” half of the finding comes from only three stores (Bethesda, Coral Gables and Sarasota), a limitation we return to in Section 8.")

# ================================================================= 7 MODEL 3
H1('7. Model 3: The Stepwise Model')
H2('7.1 Selection')
P("To let the data choose the final variables, we ran stepwise regression in JASP with the eight original variables and the 16 interactions as candidates (24 in all), with entry at p < 0.05 and removal at p > 0.10. At each step JASP adds the candidate with the largest partial correlation, provided its partial F-test is significant, and then removes any variable whose p-value has risen above 0.10. The procedure stopped after six steps. Natural, Fitness, plain Endcap and the other 15 interactions never entered, and nothing was removed. In particular, Endcap \u00d7 Sales Rep entered first and plain Endcap never did, which means the endcap effect sits entirely inside the interaction.")
labs = {'Rep_x_Endcap': 'Endcap × Sales Rep', 'Rep': 'Sales Rep', 'D13': 'Demo1-3', 'Demo': 'Demo', 'D45': 'Demo4-5', 'Price': 'Price'}
rows = [[f"Step {x['step']}", labs[x['var']], n1(x['fchg']), n3(x['r2']), n3(x['ar2']), n3(x['dr2'])] for x in R['stepwise']['path']]
table(['Step', 'Variable entered', 'F change', 'R²', 'Adj. R²', 'R² change'], rows, [0.8, 2.0, 1.0, 0.8, 0.9, 1.0], ['l', 'l', 'r', 'r', 'r', 'r'], cap='Table 7.1: Stepwise selection of Model 3')
H2('7.2 The model')
rows = []
for v, lab in [('Intercept', 'Intercept'), ('Rep_x_Endcap', 'Endcap × Sales Rep'), ('Rep', 'Sales Rep'), ('Demo', 'Demo'), ('D13', 'Demo1-3'), ('D45', 'Demo4-5'), ('Price', 'Price')]:
    c = C3[v]; rows.append([lab, neg(c['b'], n2), n2(c['se']), '' if v == 'Intercept' else neg(m2['std_beta'][v], n3), neg(c['t'], n2), p_(c['p']), f"{neg(c['lo'])} to {neg(c['hi'])}", '' if v == 'Intercept' else n2(m2['vif'][v])])
table(['Variable', 'Coefficient', 'Std. error', 'Std. beta', 't', 'p-value', '95% interval', 'VIF'], rows, [1.45, 0.85, 0.75, 0.7, 0.65, 0.7, 1.2, 0.5], size=9, cap='Table 7.2: Model 3 coefficients (dependent variable: Units)')
P(f"Units = {n2(C3['Intercept']['b'])} + {n2(C3['Rep_x_Endcap']['b'])} (Endcap × Sales Rep) + {n2(C3['Rep']['b'])} Sales Rep + {n2(C3['Demo']['b'])} Demo + {n2(C3['D13']['b'])} Demo1-3 + {n2(C3['D45']['b'])} Demo4-5 {MINUS} {n2(-C3['Price']['b'])} Price", 10.5, 6, 'c', italic=True)
P(f"Model 3 is significant (F = {n1(m2['F'])}, p < 0.001) and explains {m2['r2'] * 100:.1f}% of the variation in weekly sales, up from {m1['r2'] * 100:.1f}% for Model 1. The standard error falls from {n1(m1['rmse'])} to {n1(m2['rmse'])} units, and all six coefficients are significant at the 1% level. The ranking by standardised beta is Endcap × Sales Rep ({n2(m2['std_beta']['Rep_x_Endcap'])}), Sales Rep ({n2(m2['std_beta']['Rep'])}), Demo1-3 ({n2(m2['std_beta']['D13'])}), Demo ({n2(m2['std_beta']['Demo'])}), Demo4-5 ({n2(m2['std_beta']['D45'])}) and Price ({neg(m2['std_beta']['Price'], n2)}).")
H2('7.3 Assumption checks for Model 3')
figure('diag_final.png', 5.6, 'Figure 7.1: Residual diagnostics for Model 3')
P(f"The two separated groups of residuals have disappeared. The residuals now form an even band around zero at every predicted value, including the endcap weeks, which sit at predicted values of roughly 680 to 960 and are scattered above and below the line like the rest. The residual standard deviation is {n0(R['sd_by_endcap']['m2'][0])} for weeks without an endcap and {n0(R['sd_by_endcap']['m2'][1])} for weeks with one (it was {n0(R['sd_by_endcap']['m1'][0])} and {n0(R['sd_by_endcap']['m1'][1])} in Model 1), so linearity and constant variance are **satisfied**. The Durbin-Watson statistic is {n2(m2['dw'])}, so independence is **satisfied** on this evidence; the data are still repeated weekly readings of the same stores, so this is a reading of the diagnostics and not a formal proof. The P-P and Q-Q plots follow the diagonal along the whole range, the standardised residuals run from {neg(R['sr_range_final'][0], n1)} to +{n1(R['sr_range_final'][1])} (against {neg(m1['sr_range'][0], n1)} in Model 1), skewness is {neg(m2['skew'], n2)} and kurtosis is {n2(m2['kurt'])}. Only {m2['n_sr3']} residuals exceed ±3 standard deviations, close to what chance gives, and none of them is an endcap week. Normality is **satisfied**. The largest Cook’s distance is {n3(m2['max_cook'])}, so no single observation drives the result.")
table(['Check', 'Model 1', 'Model 2 (log)', 'Model 3'],
      [['Linearity and constant variance', 'Fails', 'Fails', 'Satisfied'], ['Durbin-Watson', n2(m1['dw']), n2(ml['dw']), n2(m2['dw'])], ['Skewness / kurtosis of residuals', f"{neg(m1['skew'], n2)} / {n1(m1['kurt'])}", f"{neg(ml['skew'], n2)} / {n1(ml['kurt'])}", f"{neg(m2['skew'], n2)} / {n2(m2['kurt'])}"],
       ['Residuals beyond ±3 SD', m1['n_sr3'], ml['n_sr3'], m2['n_sr3']], ['Largest VIF', n2(max(m1['vif'].values())), n2(max(ml['vif'].values())), n2(max(m2['vif'].values()))],
       ['R²', n3(m1['r2']), n3(ml['r2']) + ' (on ln Units)', n3(m2['r2'])]],
      [2.6, 1.2, 1.5, 1.1], ['l', 'r', 'r', 'r'], cap='Table 7.3: The three models compared')
H2('7.4 Robustness')
P(f"We ran four further checks. First, the same stores appear in every week, so we re-estimated Model 3 with standard errors clustered by store; every coefficient stays significant at p < 0.001 (for Endcap × Sales Rep the standard error is {n1(m2['cluster']['Rep_x_Endcap']['se'])} against {n1(C3['Rep_x_Endcap']['se'])}, for Demo {n1(m2['cluster']['Demo']['se'])} against {n1(C3['Demo']['se'])}). Second, in ten-fold cross-validation that holds out whole stores, the prediction error (RMSE) is {n1(R['cv_rmse']['final'])} units for Model 3 against {n1(R['cv_rmse']['base'])} for Model 1, so the gain is not over-fitting. Third, when we repeated the stepwise search on every two-, three- and four-way product of the eight variables (154 terms, 108 with enough non-zero rows), the same six terms were selected each time. Fourth, adding each of the other 15 interactions to Model 3 leaves every one insignificant (p from {n2(min(R['q5']['other15'].values()))} to {n2(max(R['q5']['other15'].values()))}). We also ran the stepwise search on ln Units; it needed {R['stepwise_ln']['n_terms']} terms and still gave R² = {n2(R['stepwise_ln']['r2'])} with kurtosis {n1(R['stepwise_ln']['kurt'])}, which is why we keep Model 3 in the original units.")

# ================================================================= 8 ANSWERS
H1('8. Answers to the Case Questions')
H2('8.1 Question 1: The model and what it shows')
P(f"Our answer is Model 3 (Table 7.2), which explains {m2['r2'] * 100:.1f}% of the variation in weekly store sales with six terms, all significant at the 1% level. Each coefficient is the change in weekly units per store when the other variables are held constant.", keep=True)
table(['Variable', 'Coefficient', '95% interval', 'Meaning'],
      [['Endcap × Sales Rep', '+' + n1(C3['Rep_x_Endcap']['b']), f"{n1(C3['Rep_x_Endcap']['lo'])} to {n1(C3['Rep_x_Endcap']['hi'])}", 'Endcap in a store with a regional rep'],
       ['Demo', '+' + n1(C3['Demo']['b']), f"{n1(C3['Demo']['lo'])} to {n1(C3['Demo']['hi'])}", 'Week of the demo'],
       ['Demo4-5', '+' + n1(C3['D45']['b']), f"{n1(C3['D45']['lo'])} to {n1(C3['D45']['hi'])}", '4 to 5 weeks after a demo'],
       ['Demo1-3', '+' + n1(C3['D13']['b']), f"{n1(C3['D13']['lo'])} to {n1(C3['D13']['hi'])}", '1 to 3 weeks after a demo'],
       ['Sales Rep', '+' + n1(C3['Rep']['b']), f"{n1(C3['Rep']['lo'])} to {n1(C3['Rep']['hi'])}", 'Regional rep, no promotion'],
       ['Price', neg(C3['Price']['b']), f"{neg(C3['Price']['lo'])} to {neg(C3['Price']['hi'])}", 'Per $1 increase in price']],
      [1.6, 1.0, 1.4, 2.4], ['l', 'r', 'r', 'l'], cap='Table 8.1: Interpretation of Model 3')
P(f"Four observations follow. Both promotional programmes go with higher sales, and the endcap effect is the largest in the model. The endcap effect is conditional: it appears only where a regional sales rep is present, and JASP did not select the plain Endcap variable once the interaction was available. The demo effect persists well beyond the demo week. And the local environment does not matter: the number of natural retailers and fitness centres nearby has no significant effect, while a regional rep and a lower price both raise sales. At the mean price and mean sales, the price elasticity is about {n2(R['elasticity'])}, so demand is inelastic over the observed range.")
H2('8.2 Question 2: Do the model’s assumptions hold?')
P(f"Not at first. Model 1 failed linearity, constant variance, independence and normality (Section 3), and the log-level Model 2 did not cure them (Section 5). Model 3 satisfies all four (Section 7.3), as Table 7.3 shows side by side. Multicollinearity is not a concern in any of the three models; the highest VIF in Model 3 is {n2(max(m2['vif'].values()))}, for Sales Rep. Beyond the plots and the Durbin-Watson statistic, we verified the model with Cook’s distance, store-clustered standard errors, store-level cross-validation and a search over all two-, three- and four-way interactions (Section 7.4). Further work that would verify it more fully: residual plots in week order for each store, a mixed model with a store effect for the repeated observations, and a plot of residuals against Price to confirm that its effect is linear (adding Price squared to Model 3 is not significant, p = {n2(R['q5']['price_sq']['p'])}).")
H2('8.3 Question 3: Do demos boost sales, and for how long?')
P(f"Yes. A demo is associated with about {n0(C3['Demo']['b'])} additional units in the week it is held. A store that has no demo effect and no endcap sells about {n0(R['baseline_no_demo_no_endcap'])} units in a week, so this is a lift of roughly {C3['Demo']['b'] / R['baseline_no_demo_no_endcap'] * 100:.0f}%. After the demo week the lift settles at about {n0(C3['D13']['b'])} to {n0(C3['D45']['b'])} units a week (roughly {C3['D13']['b'] / R['baseline_no_demo_no_endcap'] * 100:.0f}%) and stays there. The coefficients for weeks 1 to 3 and weeks 4 to 5 are almost identical (a partial F-test of equality gives F = {n2(R['q3']['D13_eq_D45']['F'])}, p = {n2(R['q3']['D13_eq_D45']['p'])}; the difference is {n1(R['q3']['diff_D45_D13']['est'])} units, 95% interval {neg(R['q3']['diff_D45_D13']['lo'], n0)} to +{n0(R['q3']['diff_D45_D13']['hi'])}), so we see no sign of the lift fading within five weeks. The demo week itself is higher than the weeks that follow (F = {n1(R['q3']['Demo_eq_D13']['F'])}, p < 0.001).")
b = R['baseline_no_demo_no_endcap']; d0, d1, d2 = C3['Demo']['b'], C3['D13']['b'], C3['D45']['b']; cum = [d0, d0 + d1, d0 + 2 * d1, d0 + 3 * d1, d0 + 3 * d1 + d2, d0 + 3 * d1 + 2 * d2]; tot = cum[-1]
lifts = [d0, d1, d1, d1, d2, d2]; names = ['Demo week', 'Week 1', 'Week 2', 'Week 3', 'Week 4', 'Week 5']
table(['Week', 'Lift (units)', 'Lift vs. baseline', 'Cumulative units', 'Share of total'], [[names[i], n1(lifts[i]), f"{lifts[i] / b * 100:.0f}%", n1(cum[i]), f"{cum[i] / tot * 100:.0f}%"] for i in range(6)], [1.4, 1.1, 1.3, 1.4, 1.2],
      cap='Table 8.2: Cumulative sales lift from one demo', size=9.5)
P(f"The baseline is {n0(b)} units, the mean weekly sales of a store with no demo effect and no endcap.", 9.5, 4, italic=True, align='l')
P(f"The lift therefore lasts at least five weeks. We cannot say when it ends, because Demo4-5 means four or more weeks after a demo and the data do not follow stores further. Adding up the weekly effects gives about {n0(R['q3']['total'])} additional units per demo over six weeks ({n1(d0)} + 3 × {n1(d1)} + 2 × {n1(d2)}), and less than a quarter of that arrives in the demo week itself. Judging demos by demo-week sales alone would understate their effect by a factor of more than four, which speaks directly to management’s worry that any increase is only temporary.")
figure('levers.png', 5.7, 'Figure 8.1: Effect of each lever in Model 3 with 95% intervals (left) and the demo lift by period (right)')
H2('8.4 Question 4: Does placement within the store affect sales?')
P(f"Yes, but only together with a regional sales rep. An endcap in a store with a rep is associated with about {n0(C3['Rep_x_Endcap']['b'])} additional units a week (95% interval {n0(C3['Rep_x_Endcap']['lo'])} to {n0(C3['Rep_x_Endcap']['hi'])}), which more than doubles that store’s sales: {n0(R['groups']['01']['mean'])} units without an endcap and {n0(R['groups']['11']['mean'])} with one. In stores served only by the national rep it makes almost no difference ({n0(R['groups']['00']['mean'])} against {n0(R['groups']['10']['mean'])} units). When plain Endcap is added back to Model 3 its coefficient is {n1(R['q4']['plain_endcap']['b'])} (p = {n2(R['q4']['plain_endcap']['p'])}).", keep=True)
table(['', 'No endcap', 'Endcap', 'Difference'], [['No regional sales rep', f"{n1(R['groups']['00']['mean'])} (n = {R['groups']['00']['n']})", f"{n1(R['groups']['10']['mean'])} (n = {R['groups']['10']['n']})", '+' + n1(R['groups']['10']['mean'] - R['groups']['00']['mean'])],
       ['Regional sales rep', f"{n1(R['groups']['01']['mean'])} (n = {R['groups']['01']['n']})", f"{n1(R['groups']['11']['mean'])} (n = {R['groups']['11']['n']})", '+' + n1(R['groups']['11']['mean'] - R['groups']['01']['mean'])]],
      [2.0, 1.6, 1.6, 1.2], cap='Table 8.3: Mean weekly units by sales rep coverage and endcap status')
P(f"Placement on its own is therefore not what drives sales. Models 1 and 2 reported a single endcap effect of about +{n0(C1['Endcap']['b'])} units, but that figure is an average of a very large effect and almost none ({n0(C3['Rep_x_Endcap']['b'])} × 36/53 = {n0(R['blend_check']['expected'])}), and it is right for neither group. Placement does not interact with demos or with price either (p = {n2(R['q5']['other15']['Endcap_x_Demo'])} for Endcap \u00d7 Demo, {n2(R['q5']['other15']['Endcap_x_D13'])} for Endcap \u00d7 Demo1-3 and {n2(R['q5']['other15']['Price_x_Endcap'])} for Price \u00d7 Endcap). One caution: the result for stores without a rep comes from three stores and {R['groups']['10']['n']} store-weeks, so we treat it as a strong indication and not a settled fact.")
H2('8.5 Question 5: Suggestions to improve the model')
P(f"Within the data we tested the obvious refinements, and none improves Model 3 (Table 8.4). The partial F-tests compare Model 3 with the same model plus the extra terms.", keep=True)
table(['Refinement added to Model 3', 'Partial F', 'p-value', 'Adj. R²'],
      [['Natural and Fitness', n2(R['q5']['natural_fitness']['F']), n2(R['q5']['natural_fitness']['p']), n3(R['q5']['natural_fitness']['ar2'])], ['Region dummies', n2(R['q5']['region']['F']), n2(R['q5']['region']['p']), n3(R['q5']['region']['ar2'])],
       ['Price squared', n2(R['q5']['price_sq']['F']), n2(R['q5']['price_sq']['p']), n3(R['q5']['price_sq']['ar2'])], ['Week trend', n2(R['q5']['week']['F']), n2(R['q5']['week']['p']), n3(R['q5']['week']['ar2'])],
       ['Store dummies (fixed effects)', n2(R['q5']['store']['F']), n2(R['q5']['store']['p']), n3(R['q5']['store']['ar2'])]],
      [3.2, 1.0, 1.0, 1.0], cap='Table 8.4: Refinements tested (Model 3 adjusted R² = ' + n3(m2['ar2']) + ')')
P("Improvements that need new data or a different method:", after=2, keep=True)
bullets([
    "Add store traffic and size. The case says demos work only in stores with enough foot traffic, and this is the most likely omitted variable.",
    "Use store and week identifiers in a mixed model with a store effect, which would deal directly with the repeated observations, and control for seasonality across May to July.",
    "Extend the demo lags beyond five weeks and split the bands into single weeks, to find where the lift actually ends.",
    f"Allow for how stores were chosen. Demos and endcaps were not assigned at random: {R['demos_in_rep_stores'][0]} of the {R['demos_in_rep_stores'][1]} demo weeks were in stores with a regional rep. If promotions go to stores that already sell well, the coefficients overstate their effect, and a planned test with comparable control stores would settle this.",
    "Test more endcaps in stores without a regional rep, so that the no-rep result no longer rests on three stores.",
    "Add cost and margin data for a demo, an endcap and a regional rep. Without them the model measures lift in units but cannot measure return on spend."], size=11)
H2('8.6 Question 6: Managerial implications and recommendations')
P("Our analysis answers two of management’s three doubts. Demos do boost sales, and the increase is not temporary. The third doubt, whether the lift justifies the cost, cannot be answered from this data because it contains no costs. With that in mind we recommend the following.", keep=True)
bullets([
    f"**Keep the demo programme.** Each demo is associated with roughly {n0(R['q3']['total'])} additional units over six weeks, and we see no sign of the lift fading in that time.",
    "**Evaluate demos over six weeks, not one.** About three quarters of the lift arrives after the demo week, so any cost-benefit calculation should use the full six-week figure.",
    "**Run endcaps only where a regional sales rep is present.** An endcap with a rep is the single most effective action in the data, while an endcap without one shows almost no return. The endcap competition should concentrate on stores with face-to-face rep coverage.",
    f"**Treat regional rep coverage as an investment.** A rep is worth about {n0(C3['Rep']['b'])} units a week alone and is also what makes an endcap work. Extending coverage to stores that are candidates for endcaps is likely to pay better than adding endcaps in stores without coverage.",
    f"**Do not rely on price cuts for volume.** With an elasticity of about {n2(R['elasticity'])}, units rise less than proportionally when the price falls, so revenue falls.",
    f"**Close the cost gap before the budget decision.** A demo pays back if it costs less than {n0(R['q3']['total'])} units times the margin per unit. Marketing should provide the cost of a demo, an endcap and a rep so that this can be worked out.",
    "**Confirm with a controlled test.** Our results are associations from 11 weeks of data in which promotions were not randomly assigned. A small planned test, especially of endcaps in more stores without a regional rep, would confirm them before the budget is reallocated."], style='List Number', size=11)

# ================================================================= 9 LEARNINGS
H1('9. Key Learnings')
bullets([
    "**A high R² does not make a model correct.** Model 1 explained 67% of the variation and almost every coefficient was significant, but it failed the assumptions. We learned to read the residuals before trusting the fit.",
    "**Extreme residuals are information, not noise.** The 23 residuals beyond three standard deviations looked like outliers. Deleting them would have discarded the most important finding in the case; asking what they had in common revealed the missing interaction.",
    "**One cause can break several assumptions at once.** Linearity, constant variance, independence and normality all failed for the same reason, and one added term fixed them together. A transformation treated only the symptoms.",
    "**An average effect can be true of nobody.** The single endcap coefficient of about 305 units averaged a very large effect with almost none.",
    "**Automatic selection needs the right candidates.** Stepwise regression only chooses from the variables it is offered. It found Endcap × Sales Rep because we put it on the list, so judgement still has to come from the analyst.",
    "**Statistical significance is not the business answer.** The model says how many units each programme adds. The budget decision also needs costs and margins, and a controlled test before the associations can be called cause and effect."], size=11)

# ================================================================= APPENDIX
H1('Appendix A: JASP Settings Used at Each Step')
table(['Step', 'JASP analysis and settings', 'Used in'],
      [['Model 1', 'Regression → Classical → Linear Regression. Dependent variable: Units. Covariates: Price, Rep, Endcap, Demo, D13, D45, Natural, Fitness. Method: Enter. Statistics: estimates, confidence intervals, model fit, R-squared change, collinearity (VIF), Durbin-Watson.', 'Tables 2.1, 2.2'],
       ['Assumptions', 'Same analysis. Residuals: statistics and casewise diagnostics (standardised residual > 3, Cook’s distance). Plots: residuals vs. predicted, residuals histogram, Q-Q plot of standardised residuals.', 'Figure 3.1, Table 3.1'],
       ['Model 2', 'New linear regression with lnUnits as the dependent variable (column in the data file) and the same eight covariates and options.', 'Table 4.1, Figure 5.1'],
       ['Interaction screen', 'Model 1 covariates plus one product column at a time (e.g. Rep_x_Endcap, the Endcap × Sales Rep term); read the R-squared change and the t-test of the added term.', 'Table 6.2'],
       ['Model 3', 'New linear regression, dependent variable Units, covariates: the eight variables and the 16 interaction columns. Method: Stepwise, with p-value criteria: entry 0.05, removal 0.10. Same statistics and plots.', 'Tables 7.1, 7.2, Figure 7.1'],
       ['Demo duration test', 'Dependent: Units; covariates Rep_x_Endcap, Rep, Demo, D13_plus_D45, Price (Demo1-3 and Demo4-5 forced equal). F = (SSE reduced − SSE full) / (SSE full / 1,379).', 'Section 8.3']],
      [1.2, 4.2, 1.0], ['l', 'l', 'l'], size=9, cap='Table A.1: JASP settings')
P("The residual figures were redrawn from the same residuals for a uniform layout, with a P-P plot added as the brief asks; JASP\u2019s own Q-Q plot shows the same information. The clustered standard errors, cross-validation, Breusch-Pagan test (p = " + f"{n3(m2['bp_p'])}" + " for Model 3) and the search over 154 higher-order terms were run outside JASP. Scripts and a workbook of all tables are in the submission folder.", 10)

out = os.path.join(HERE, 'GoodBelly_Case_Report_Sec_E_Group_15.docx'); doc.save(out); print('saved', out)
